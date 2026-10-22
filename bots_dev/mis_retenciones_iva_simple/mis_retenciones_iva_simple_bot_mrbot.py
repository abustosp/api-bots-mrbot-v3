import asyncio
import os
import re
import tempfile
import time
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, List, Optional, Sequence

from minio.error import S3Error
from playwright.async_api import async_playwright

from shared.arca_login import login_arca, open_arca_service
from app.utils.bucket import subir_archivo_a_minio
from shared.proxy import get_proxy_config
from app.utils.browser_limiter import limit_browser

# Constants and Helpers
INVALID_FILENAME_RE = re.compile(r"[\\/:*?\"<>|]+")

def safe_filename(value: str) -> str:
    cleaned = INVALID_FILENAME_RE.sub("_", value or "")
    cleaned = re.sub(r"\s+", "_", cleaned).strip("_")
    return cleaned or "download"

def _sanitize_component(value: str) -> str:
    cleaned = INVALID_FILENAME_RE.sub("_", value or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned

def _format_date_component(value: str) -> str:
    return re.sub(r"\s+", " ", (value or "").strip()).replace("/", "-")

def _tipo_to_abbr(tipo: str) -> str:
    lowered = (tipo or "").strip().lower()
    if lowered.startswith("ret"):
        return "ret"
    if lowered.startswith("per"):
        return "per"
    return _sanitize_component(tipo)

def build_retenciones_filename(
    cuit_representado: str,
    desde: str,
    hasta: str,
    impuesto_label: str,
    tipo: str,
    denominacion: str,
) -> str:
    cuit_digits = re.sub(r"\D", "", cuit_representado or "")
    fin_cuit = cuit_digits[-1:] if cuit_digits else ""
    cuit_value = cuit_digits or (cuit_representado or "")
    parts = [
        fin_cuit,
        cuit_value,
        _format_date_component(desde),
        _format_date_component(hasta),
        impuesto_label,
        _tipo_to_abbr(tipo),
        denominacion,
    ]
    cleaned_parts = [_sanitize_component(part) for part in parts if part]
    nombre = " - ".join(cleaned_parts).strip()
    if not nombre.endswith(".csv"):
        nombre = f"{nombre}.csv" if nombre else "reporte.csv"
    return nombre

minio_url_interna = os.getenv("MINIO_URL_INTERNA")
minio_url_externa = os.getenv("MINIO_URL_EXTERNA")
minio_ssl = os.getenv("MINIO_SSL", "False").lower() == "true"
minio_interno = os.getenv("MINIO_INTERNO", "False").lower() == "true"

minio_bucket_iva_simple = os.getenv("MINIO_BUCKET_MISRETENCIONESIVASIMPLE", "misretencionesivasimple")
minio_access_key_iva_simple = os.getenv("MINIO_MRBOT_ACCESS_KEY")
minio_secret_key_iva_simple = os.getenv("MINIO_MRBOT_SECRET_KEY")

def _unique(values: Sequence[str]) -> List[str]:
    seen = set()
    ordered: List[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            ordered.append(value)
    return ordered

def _impuesto_variants(label: str, code: str) -> List[str]:
    normalized = re.sub(r"\s+", " ", (label or "").strip())
    variants: List[str] = [normalized] if normalized else []
    text_only = normalized
    if code:
        text_only = re.sub(rf"^{re.escape(code)}\s*-\s*", "", normalized, flags=re.I).strip()
        if normalized and not normalized.startswith(code):
            variants.append(f"{code} - {normalized}")
    if text_only:
        variants.append(text_only)
        variants.append(f"- {text_only}")
    no_punct = re.sub(r"[./]", " ", normalized)
    if no_punct:
        variants.append(no_punct)
    if text_only:
        text_only_no_punct = re.sub(r"[./]", " ", text_only)
        variants.append(text_only_no_punct)
        variants.append(f"- {text_only_no_punct}")
    if normalized:
        variants.extend([normalized.upper(), normalized.lower()])
    if text_only:
        variants.extend([text_only.upper(), text_only.lower()])
    cleaned = [re.sub(r"\s+", " ", value).strip() for value in variants if value]
    return _unique(cleaned)

def _impuesto_patterns(label: str, code: str) -> List[re.Pattern]:
    normalized = re.sub(r"\s+", " ", (label or "").strip())
    if code:
        normalized = re.sub(rf"^{re.escape(code)}\s*-\s*", "", normalized, flags=re.I).strip()
    normalized = re.sub(r"[\s\-./,]+", " ", normalized).strip()
    tokens = [re.escape(token) for token in normalized.split(" ") if token]
    if not tokens:
        return []
    flex = r"[\s\-./,]+".join(tokens)
    patterns = [re.compile(flex, re.I)]
    if code:
        patterns.insert(0, re.compile(rf"\b{re.escape(code)}\b[\s\-./,]*{flex}", re.I))
    return patterns

def format_cuit(cuit: str) -> str:
    digits = re.sub(r"\D", "", cuit or "")
    if len(digits) == 11:
        return f"{digits[:2]}-{digits[2:10]}-{digits[10:]}"
    return cuit or ""

def cuit_regex(cuit: str) -> re.Pattern:
    digits = re.sub(r"\D", "", cuit or "")
    if not digits:
        return re.compile(re.escape(cuit or ""))
    if len(digits) == 11:
        pattern = rf"{digits[:2]}\D*{digits[2:10]}\D*{digits[10:]}"
    else:
        pattern = r"\D*".join(list(digits))
    return re.compile(pattern)

IVA_SIMPLE_RESULT_LABEL = "IMPUESTOS IVA SIMPLE"
IVA_SIMPLE_OPERACIONES = (
    "RETENCIONES IMPOSITIVAS",
    "PERCEPCIONES IMPOSITIVAS",
    "PERCEPCIONES ADUANERAS",
)


def validate_iva_simple_date_range(desde: str, hasta: str) -> Optional[str]:
    """Valida que hasta no supere el último día del mes siguiente a desde."""
    try:
        desde_dt = datetime.strptime((desde or "").strip(), "%d/%m/%Y")
        hasta_dt = datetime.strptime((hasta or "").strip(), "%d/%m/%Y")
    except ValueError:
        return "Las fechas deben tener formato DD/MM/AAAA"
    if hasta_dt < desde_dt:
        return "La fecha hasta no puede ser anterior a la fecha desde"
    if desde_dt.month == 12:
        next_month = datetime(desde_dt.year + 1, 1, 1)
    else:
        next_month = datetime(desde_dt.year, desde_dt.month + 1, 1)
    max_hasta = next_month + timedelta(days=31)
    max_hasta = max_hasta.replace(day=1) - timedelta(days=1)
    if hasta_dt > max_hasta:
        return (
            f"La fecha hasta no puede ser posterior a {max_hasta.strftime('%d/%m/%Y')} "
            f"para la fecha desde {desde_dt.strftime('%d/%m/%Y')}"
        )
    return None


async def set_iva_simple_mode(page) -> bool:
    for value in ("siap", "ivaSimple"):
        checkbox = page.locator(f"input[type='checkbox'][value='{value}']")
        try:
            if await checkbox.count() == 0:
                return False
            if value == "ivaSimple":
                if not await checkbox.is_checked():
                    await checkbox.check(timeout=5000)
            elif await checkbox.is_checked():
                await checkbox.uncheck(timeout=5000)
        except Exception:
            return False
    return True


async def select_iva_simple_operation(page, operation: str) -> bool:
    selector = page.locator("#selectTipoOperacionIvaSimple")
    try:
        await selector.wait_for(state="visible", timeout=5000)
        await selector.click()
    except Exception:
        return False
    candidates = [
        page.get_by_role("option", name=re.compile(rf"^{re.escape(operation)}$", re.I)),
        page.locator("#selectTipoOperacionIvaSimple-dropdown").get_by_text(operation, exact=True),
    ]
    for candidate in candidates:
        try:
            if await candidate.count():
                await candidate.first.click(timeout=3000)
                return True
        except Exception:
            continue
    return False


async def ensure_iva_simple_date_range(page, desde: str, hasta: str) -> bool:
    if validate_iva_simple_date_range(desde, hasta):
        return False
    fields = page.locator("input[placeholder='dd/mm/aaaa']")
    try:
        await fields.first.wait_for(state="visible", timeout=5000)
        if await fields.count() < 2:
            return False
        for field, value in ((fields.nth(0), desde), (fields.nth(1), hasta)):
            await field.fill(value)
            await field.press("Enter")
        return True
    except Exception:
        return False

def _text_has_cuit(text: str, formatted: str, digits_only: str) -> bool:
    if not text:
        return False
    cleaned = text.replace("[", "").replace("]", "")
    return (formatted and formatted in cleaned) or (digits_only and digits_only in cleaned)

async def verify_representado_selected(
    page,
    cuit_representado: str,
    wait_timeout_ms: int = 4000,
) -> bool:
    formatted = format_cuit(cuit_representado)
    digits_only = re.sub(r"\D", "", cuit_representado or "")
    if not formatted and not digits_only:
        return False

    navbar = page.locator("#navBarMisRetenciones")
    try:
        await navbar.wait_for(state="visible", timeout=wait_timeout_ms)
    except Exception:
        pass

    cuit_locator = page.locator("#navBarMisRetenciones-relationCuil")
    label_locator = page.locator("#navBarMisRetenciones-relationLabel")
    for locator in (cuit_locator, label_locator):
        try:
            if await locator.count():
                text_content = await locator.first.text_content()
                if _text_has_cuit(text_content or "", formatted, digits_only):
                    return True
        except Exception:
            continue

    try:
        nav_text = await navbar.text_content()
        if _text_has_cuit(nav_text or "", formatted, digits_only):
            return True
    except Exception:
        pass

    return False

async def select_representado(page, cuit_representado: str) -> None:
    formatted = format_cuit(cuit_representado)
    cuit_pattern = cuit_regex(cuit_representado)
    try:
        await page.locator(".e-relation__card").first.wait_for(state="visible", timeout=20000)
    except Exception:
        pass
    attempts = [
        page.locator(".e-relation__card", has=page.get_by_text(formatted)),
        page.locator(".e-relation__card", has=page.get_by_text(cuit_pattern)),
        page.locator("div[id^=selectorRelaciones_relation]", has_text=formatted),
        page.get_by_text(formatted),
        page.get_by_text(cuit_pattern),
    ]
    for locator in attempts:
        try:
            await locator.first.click(timeout=3000)
            return
        except Exception:
            continue

async def wait_for_ready(page, timeout_ms: int = 18000) -> None:
    try:
        await page.wait_for_load_state("domcontentloaded", timeout=timeout_ms)
    except Exception:
        pass
    try:
        await page.wait_for_load_state("networkidle", timeout=min(timeout_ms, 9000))
    except Exception:
        pass

async def wait_for_results_state(page, timeout_ms: int = 30000) -> str:
    deadline = time.monotonic() + (timeout_ms / 1000)
    while time.monotonic() < deadline:
        try:
            if await page.locator("#modal-sire").is_visible(timeout=500):
                return "sire"
        except Exception:
            pass

        try:
            if await page.locator("#modal-confirmar").is_visible(timeout=500):
                return "error"
        except Exception:
            pass

        try:
            if await page.locator("#modal-sinresultados").is_visible(timeout=500):
                return "export_processing"
        except Exception:
            pass

        try:
            body_text = (await page.locator("body").inner_text(timeout=500) or "").lower()
            if "no se encontraron resultados para esta consulta" in body_text:
                return "no_results"
            if "no hay resultados para tu consulta" in body_text:
                return "no_results"
        except Exception:
            pass

        export_btn = page.locator("#btnExportarOtrosFormatos")
        try:
            if await export_btn.count() > 0 and await export_btn.is_visible(timeout=500):
                return "results"
        except Exception:
            pass

        try:
            table = page.locator("#tablaResultados")
            if await table.count() > 0 and await table.is_visible(timeout=500):
                return "results"
        except Exception:
            pass

        if re.search(r"/resultados", page.url, re.I):
            await asyncio.sleep(0.35)
            if await export_btn.count() > 0:
                return "results"

        await asyncio.sleep(0.2)

    return "no_results"

async def open_representado_selector(page) -> None:
    cards = page.locator(".e-relation__card")
    if await cards.count():
        return
    try:
        await page.locator("#navBarMisRetenciones").wait_for(state="visible", timeout=20000)
    except Exception:
        pass
    relation_link = page.locator("#navBarMisRetenciones-relationLabel")
    relation_cuit = page.locator("#navBarMisRetenciones-relationCuil")
    for locator in (relation_link, relation_cuit):
        if await locator.count():
            try:
                await locator.click(timeout=3000)
                try:
                    await cards.first.wait_for(state="visible", timeout=5000)
                    return
                except Exception:
                    continue
            except Exception:
                continue
    toggle = page.locator("#e-navbar-dropdown-toggle")
    if await toggle.count():
        try:
            await toggle.wait_for(state="visible", timeout=5000)
            await toggle.click()
        except Exception:
            pass
    else:
        fallback_toggle = page.locator("#navBarMisRetenciones_avatar")
        if await fallback_toggle.count():
            await fallback_toggle.click()
        else:
            try:
                await page.locator("#navBarMisRetenciones [data-bs-toggle='dropdown']").first.click()
            except Exception:
                try:
                    await page.get_by_role("button", name=re.compile("menu|usuario|perfil", re.I)).click()
                except Exception:
                    pass
    dropdown = page.locator(".dropdown-menu.show")
    try:
        await dropdown.wait_for(state="visible", timeout=5000)
    except Exception:
        pass
    candidates = [
        page.locator("#navBarMisRetenciones-dropdown-changeRelation"),
        page.get_by_role("link", name=re.compile("Seleccionar representado", re.I)),
        page.get_by_role("link", name=re.compile("Cambiar representado", re.I)),
        page.get_by_text(re.compile("Seleccionar representado", re.I)),
        page.get_by_text(re.compile("Cambiar representado", re.I)),
    ]
    for locator in candidates:
        try:
            await locator.first.click(timeout=3000)
            try:
                await cards.first.wait_for(state="visible", timeout=5000)
            except Exception:
                pass
            return
        except Exception:
            continue

async def _get_selected_impuesto_text(page) -> str:
    candidates = [
        page.locator("#selectImpuestos-assist"),
        page.locator(".multiselect-single-label-text"),
    ]
    for locator in candidates:
        try:
            if await locator.count():
                text = (await locator.first.text_content()) or ""
                text = re.sub(r"\s+", " ", text).strip()
                if text:
                    return text
        except Exception:
            continue
    return ""

async def is_impuesto_selected(page, label: str, code: str) -> bool:
    current = await _get_selected_impuesto_text(page)
    if not current:
        return False
    if code and code in current:
        return True
    normalized_label = re.sub(r"\s+", " ", (label or "").strip())
    return bool(normalized_label and normalized_label in current)

async def _set_datepicker_value(page, input_locator, value: str, label: str) -> bool:
    """Set a datepicker value with multiple strategies: fill, type, and JS fallback.
    The input may be hidden (Angular controlled), so we don't require visibility."""
    if not value:
        return False
    for attempt in range(3):
        try:
            await input_locator.wait_for(state="attached", timeout=5000)
        except Exception:
            await asyncio.sleep(0.5)
            continue
        # Try to focus/click (may fail if hidden, that's ok)
        try:
            await input_locator.click(timeout=2000, force=True)
        except Exception:
            pass
        # Strategy 1: fill + Enter
        try:
            await input_locator.fill("")
            await input_locator.fill(value, force=True)
            await input_locator.press("Enter")
            await asyncio.sleep(0.3)
            current = (await input_locator.input_value()) or ""
            if current.strip() == value:
                return True
        except Exception:
            pass
        # Strategy 2: type char by char + Enter
        if attempt == 1:
            try:
                await input_locator.click(timeout=2000, force=True)
                await input_locator.fill("", force=True)
                await input_locator.type(value, delay=50)
                await input_locator.press("Enter")
                await asyncio.sleep(0.3)
                current = (await input_locator.input_value()) or ""
                if current.strip() == value:
                    return True
            except Exception:
                pass
        # Strategy 3: JS value injection + dispatch events
        if attempt == 2:
            try:
                el = await input_locator.element_handle()
                if el:
                    await el.evaluate(
                        "(el, val) => { el.value = val; "
                        "el.dispatchEvent(new Event('input', {bubbles:true})); "
                        "el.dispatchEvent(new Event('change', {bubbles:true})); "
                        "el.dispatchEvent(new Event('blur', {bubbles:true})); }",
                        value,
                    )
                    await asyncio.sleep(0.3)
                    current = (await input_locator.input_value()) or ""
                    if current.strip() == value:
                        return True
            except Exception:
                pass
        await asyncio.sleep(0.4)
    # Final check: even if we can't read the value, the fill may have worked
    # (Angular-controlled inputs sometimes don't expose value via input_value())
    try:
        current = (await input_locator.input_value()) or ""
        if current.strip():
            return True
    except Exception:
        pass
    return False

async def ensure_date_range(page, desde: str, hasta: str) -> bool:
    if not desde or not hasta:
        return False
    # Datepickers reales de la pantalla de consulta (Siap pair deprecated: ya no existe)
    datepicker_pairs = [
        ("#fechaRetencionDesde__input", "#fechaRetencionHasta__input"),
    ]
    for desde_sel, hasta_sel in datepicker_pairs:
        fecha_desde_input = page.locator(desde_sel)
        fecha_hasta_input = page.locator(hasta_sel)
        try:
            await fecha_desde_input.wait_for(state="attached", timeout=5000)
            await fecha_hasta_input.wait_for(state="attached", timeout=5000)
        except Exception:
            continue
        try:
            current_desde = (await fecha_desde_input.input_value()) or ""
            current_hasta = (await fecha_hasta_input.input_value()) or ""
        except Exception:
            current_desde = ""
            current_hasta = ""
        if current_desde.strip() == desde and current_hasta.strip() == hasta:
            return True
        if not await _set_datepicker_value(page, fecha_desde_input, desde, "desde"):
            continue
        await asyncio.sleep(0.3)
        if not await _set_datepicker_value(page, fecha_hasta_input, hasta, "hasta"):
            continue
        await asyncio.sleep(0.3)
        # Final verification
        try:
            final_desde = (await fecha_desde_input.input_value()) or ""
            final_hasta = (await fecha_hasta_input.input_value()) or ""
            if final_desde.strip() == desde and final_hasta.strip() == hasta:
                return True
            # If both have some value, consider it success (Angular may reformat)
            if final_desde.strip() and final_hasta.strip():
                return True
        except Exception:
            return True
    return False


async def set_export_mode(page, mode: Optional[str]) -> bool:
    expected = (mode or "").strip().lower()
    if expected not in {"", "siap"}:
        return False
    for value in ("ivaSimple", "siap"):
        checkbox = page.locator(f"input[type='checkbox'][value='{value}']")
        try:
            if await checkbox.count() == 0:
                continue
            checked = await checkbox.is_checked()
            should_check = value.lower() == expected
            if should_check and not checked:
                await checkbox.check(timeout=5000)
            elif not should_check and checked:
                await checkbox.uncheck(timeout=5000)
        except Exception:
            return False
    return True


async def ensure_inicio(
    page,
    cuit_representado: str,
    timeout_ms: int = 20000,
) -> bool:
    async def _click_nueva_consulta() -> None:
        nueva_tab = page.locator("#tabNuevaConsulta-tab")
        if await nueva_tab.count():
            try:
                selected = await nueva_tab.get_attribute("aria-selected")
                if selected != "true":
                    await nueva_tab.click(timeout=5000)
            except Exception:
                pass

    async def _ensure_relation() -> bool:
        if cuit_representado:
            if not await verify_representado_selected(page, cuit_representado):
                await open_representado_selector(page)
                await select_representado(page, cuit_representado)
                return await verify_representado_selected(page, cuit_representado)
        return True

    # Fast path: el formulario ya esta disponible, no navegar (un reload resetea la relacion).
    await _click_nueva_consulta()
    representado_ok = await _ensure_relation()
    impuesto_input = page.locator("#selectImpuestos")
    try:
        await impuesto_input.wait_for(state="visible", timeout=5000)
        return representado_ok
    except Exception:
        pass

    # El formulario no esta visible: navegar a inicio (recarga completa) y rearmar todo.
    # IMPORTANTE: la ruta real del formulario (HomeView) es la raíz '/', NO '/inicio'.
    # '/inicio' no existe en el router de la SPA, por eso el formulario nunca se renderiza.
    inicio_url = "https://mirequa-web.arca.gob.ar/"
    breadcrumb = page.locator("#bcMisRetenciones-li-0-routerLink")
    if await breadcrumb.count():
        try:
            await breadcrumb.first.click(timeout=5000)
        except Exception:
            await page.goto(inicio_url, timeout=timeout_ms)
    else:
        await page.goto(inicio_url, timeout=timeout_ms)

    await _click_nueva_consulta()
    representado_ok = await _ensure_relation()
    try:
        await page.locator("#selectImpuestos").wait_for(state="visible", timeout=timeout_ms)
    except Exception:
        pass

    return representado_ok

async def open_impuesto_dropdown(page, label: str, code: str) -> None:
    input_locator = page.locator("#selectImpuestos")
    try:
        await input_locator.wait_for(state="visible", timeout=20000)
        await input_locator.click()
        try:
            await input_locator.press("Control+A")
            await input_locator.press("Backspace")
        except Exception:
            pass
    except Exception:
        try:
            await page.get_by_role("combobox", name=re.compile("Impuesto", re.I)).click()
        except Exception:
            pass
    dropdown = page.locator("#selectImpuestos-dropdown")
    try:
        if not await dropdown.is_visible():
            wrapper = page.locator("#selectImpuestos").locator("..")
            if await wrapper.count():
                await wrapper.click()
        if not await dropdown.is_visible():
            caret = page.locator("#selectImpuestos_caret")
            if await caret.count():
                await caret.click()
        if not await dropdown.is_visible():
            multiselect = page.locator(".multiselect-wrapper")
            if await multiselect.count():
                await multiselect.click()
        await dropdown.wait_for(state="visible", timeout=5000)
    except Exception:
        pass
    try:
        selected_options = page.locator("#selectImpuestos-multiselect-options li.multiselect-option.is-selected")
        selected_count = await selected_options.count()
        for idx in range(selected_count):
            option = selected_options.nth(idx)
            try:
                await option.click(timeout=2000)
            except Exception:
                pass
    except Exception:
        pass

async def select_impuesto(page, label: str, code: str) -> bool:
    if await is_impuesto_selected(page, label, code):
        return True
    await open_impuesto_dropdown(page, label, code)
    if code:
        prefixes = ("IMP", "SS", "ADU", "SIR")
        for prefix in prefixes:
            option_id = f"#selectImpuestos-multiselect-option-{prefix}_{code}"
            option = page.locator(option_id)
            if await option.count():
                try:
                    await option.first.click(timeout=2000)
                    return True
                except Exception:
                    pass
    options_root = page.locator("#selectImpuestos-dropdown")
    group = options_root.locator('[role="group"][aria-label="Impositivas"]')
    scopes = []
    try:
        if await group.count():
            scopes.append(group)
    except Exception:
        pass
    scopes.append(options_root)
    for scope in scopes:
        for candidate in _impuesto_variants(label, code):
            try:
                await scope.get_by_role("option", name=candidate).first.click(timeout=2000)
                return True
            except Exception:
                continue
        for pattern in _impuesto_patterns(label, code):
            try:
                await scope.get_by_role("option", name=pattern).first.click(timeout=2000)
                return True
            except Exception:
                continue
    return False

async def select_tipo(page, tipo: str) -> bool:
    if not tipo:
        return False
    container = page.locator("#tipoOperacion")
    try:
        await container.wait_for(state="visible", timeout=5000)
    except Exception:
        return False
    labels = container.locator("label")
    try:
        label_count = await labels.count()
    except Exception:
        label_count = 0
    if label_count == 0:
        return False
    if tipo.lower().startswith("ret"):
        pattern = re.compile("retenci[o\\u00f3]n", re.I)
    elif tipo.lower().startswith("per"):
        pattern = re.compile("percepci[o\\u00f3]n", re.I)
    else:
        pattern = re.compile(re.escape(tipo), re.I)
    try:
        target = labels.filter(has_text=pattern)
        if await target.count():
            await target.first.click(timeout=3000)
            return True
    except Exception:
        pass
    try:
        await container.get_by_label(pattern).first.check(timeout=3000)
        return True
    except Exception:
        pass
    return False

async def _click_with_retry(
    locator,
    timeout_ms: int = 5000,
    retries: int = 3,
    force: bool = False,
) -> bool:
    """Try to click a locator with retries, handling overlays and stale elements."""
    for attempt in range(retries):
        try:
            if await locator.count() == 0:
                await asyncio.sleep(0.3)
                continue
            try:
                if await locator.first.is_visible(timeout=1000):
                    await locator.first.click(timeout=timeout_ms, force=force)
                    return True
            except Exception:
                pass
            # Try force-click on later attempts
            if attempt >= 1:
                try:
                    await locator.first.click(timeout=timeout_ms, force=True)
                    return True
                except Exception:
                    pass
            await asyncio.sleep(0.4)
        except Exception:
            await asyncio.sleep(0.3)
    return False

async def select_export_option(page, export_option: str) -> bool:
    patterns = {
        "iva_simple": re.compile(r"Exportar\s+para\s+IVA\s+SIMPLE", re.I),
        "siap": re.compile(r"Exportar\s+para\s+SIAP", re.I),
        "xls": re.compile(r"Exportar\s+a\s+Excel\s*\(\s*\.xls\s*\)", re.I),
    }
    pattern = patterns.get((export_option or "").strip().lower())
    if pattern is None:
        return False
    dropdown = page.locator(".dropdown-menu.show")
    scopes = (dropdown, page)
    for scope in scopes:
        for locator in (
            scope.get_by_role("button", name=pattern),
            scope.locator("button", has_text=pattern),
        ):
            try:
                if await locator.count() and await locator.first.is_visible(timeout=3000):
                    await locator.first.click(timeout=5000)
                    return True
            except Exception:
                continue
    return False


async def download_csv(
    page,
    impuesto_code: str,
    impuesto_label: str,
    tipo: str,
    nombre_archivo: str,
    base_dir: Path,
    error_logs: list[str],
    required_result_marker: Optional[str] = None,
    export_option: str = "xls",
) -> Optional[Path]:
    safe_impuesto = safe_filename(impuesto_code or impuesto_label or "impuesto")
    safe_tipo = safe_filename(tipo or "tipo")
    contexto = " ".join([value for value in [impuesto_code, impuesto_label, tipo] if value]).strip()

    # --- Click "Consultar" with retries ---
    consultar_clicked = False
    for attempt in range(3):
        consultar_attempts = [
            page.locator("#btnConsultarRetenciones"),
            page.get_by_role("button", name=re.compile("Consultar", re.I)),
            page.locator("button:has-text('Consultar')"),
        ]
        for locator in consultar_attempts:
            if await _click_with_retry(locator, timeout_ms=5000, retries=2, force=(attempt > 0)):
                consultar_clicked = True
                break
        if consultar_clicked:
            break
        await asyncio.sleep(0.5)

    if not consultar_clicked:
        error_logs.append(f"No se pudo ejecutar la consulta para {contexto or 'Mis Retenciones'}")
        return None

    await wait_for_ready(page)
    state = await wait_for_results_state(page)
    if state == "sire":
        try:
            await page.locator("#modal-sire_btnOK").click(timeout=3000)
        except Exception:
            pass
        error_logs.append(f"El servicio SIRE requiere gestionar la consulta manualmente ({contexto})")
        return None
    if state == "error":
        try:
            await page.locator("#modal-confirmar_btnOK").click(timeout=3000)
        except Exception:
            pass
        error_logs.append(f"Se detecto un error al consultar {contexto}")
        return None
    if state == "no_results":
        error_logs.append(f"No se encontraron resultados para {contexto}")
        return None
    if state != "results":
        error_logs.append(f"No se encontraron resultados para {contexto}")
        return None

    if required_result_marker:
        body_text = (await page.locator("body").inner_text() or "")
        if required_result_marker.lower() not in body_text.lower():
            error_logs.append(
                f"No se encontró la línea esperada '{required_result_marker}' en los resultados de {contexto}"
            )
            return None

    no_results_messages = [
        page.locator("text=/no se encontraron resultados/i"),
        page.locator("text=/sin resultados/i"),
        page.locator(".alert:has-text('No se encontraron')"),
        page.locator("#mensajeSinResultados"),
        page.locator("text=/no se encontraron registros/i"),
    ]

    for no_results_locator in no_results_messages:
        try:
            if await no_results_locator.count() > 0 and await no_results_locator.is_visible(timeout=1000):
                error_logs.append(f"No se encontraron resultados para {contexto}")
                return None
        except Exception:
            pass

    # --- Click "Exportar" with retries ---
    export_clicked = False
    for attempt in range(3):
        export_button = page.locator("#btnExportarOtrosFormatos")
        if await _click_with_retry(export_button, timeout_ms=5000, retries=2, force=(attempt > 0)):
            export_clicked = True
            break
        # Fallbacks
        export_attempts = [
            page.get_by_role("button", name=re.compile("Exportar", re.I)),
            page.locator("button:has-text('Exportar')"),
        ]
        for locator in export_attempts:
            if await _click_with_retry(locator, timeout_ms=3000, retries=2, force=(attempt > 0)):
                export_clicked = True
                break
        if export_clicked:
            break
        await asyncio.sleep(0.5)

    if not export_clicked:
        error_logs.append(f"No se pudo exportar resultados para {contexto}")
        return None

    if not await select_export_option(page, export_option):
        error_logs.append(
            f"No se pudo seleccionar el formato de exportacion {export_option} para {contexto}"
        )
        return None

    export_all_ok = page.locator("#tablaResultadosexport_all_btnOK")
    try:
        if await export_all_ok.is_visible(timeout=4000):
            await export_all_ok.click(timeout=5000)
    except Exception:
        pass

    # --- Confirm download with retries ---
    exportar_archivo_clicked = False
    ver_archivo = page.locator("#modal-exportar-consulta_btnOK, #modal-sinresultados_btnOK")
    for attempt in range(3):
        try:
            await ver_archivo.first.wait_for(state="visible", timeout=10000)
            if await _click_with_retry(ver_archivo, timeout_ms=5000, retries=2, force=(attempt > 0)):
                exportar_archivo_clicked = True
                break
        except Exception:
            pass
        await asyncio.sleep(0.4)

    if not exportar_archivo_clicked:
        error_logs.append(f"No se pudo confirmar la exportacion del archivo para {contexto}")
        return None

    # --- Wait for download link (pestaña "consultas exportadas", ag-grid) con fallbacks históricos ---
    download_link = None
    link_attempts = [
        page.locator('#tablaExportarAplicativo .ag-row[row-index="0"] a[title="Descargar Archivo"]').first,
        page.locator("#tablaExportarAplicativo a[title='Descargar Archivo']").first,
        page.locator("a.boton_mt[download], a[download][href*='descarga']").first,
    ]
    for locator in link_attempts:
        try:
            await locator.wait_for(state="visible", timeout=8000)
            download_link = locator
            break
        except Exception:
            continue
    if download_link is None:
        error_logs.append(f"No se encontro el link de descarga para {contexto}")
        return None

    download = None
    for attempt in range(3):
        try:
            async with page.expect_download(timeout=20000) as download_info:
                await download_link.click(timeout=5000, force=(attempt > 0))
            download = await download_info.value
            break
        except Exception:
            await asyncio.sleep(0.5)

    if not download:
        error_logs.append(f"No se pudo descargar el archivo exportado para {contexto}")
        return None

    filename = _sanitize_component(nombre_archivo) or f"retenciones_{safe_impuesto}_{safe_tipo}"
    stem = filename[:-4] if filename.lower().endswith(".csv") else filename
    suggested_suffix = Path(download.suggested_filename or "").suffix
    filename = f"{stem}{suggested_suffix or '.csv'}"

    base_dir.mkdir(parents=True, exist_ok=True)
    download_path = base_dir / filename

    try:
        await download.save_as(str(download_path))
    except Exception as exc:
        error_logs.append(f"No se pudo guardar el CSV para {contexto}: {exc}")
        return None

    return download_path

def archivo_coincide_operacion(archivo: Path, operacion: str) -> bool:
    """Falso si ``Descripcion Operacion`` trae filas de la otra operacion.

    ARCA a veces entrega en Percepciones el mismo archivo de Retenciones. Los
    archivos sin esa columna (formato fijo iva_simple) no se pueden validar.
    """
    esperado = "PERCEP" if "PERCEP" in operacion.upper() else "RETENC"
    try:
        if archivo.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            filas = list(load_workbook(archivo, read_only=True).active.iter_rows(values_only=True))
            encabezado, datos = (filas[0] if filas else ()), filas[1:]
        else:
            import csv

            with open(archivo, encoding="utf-8-sig", newline="") as fh:
                filas = list(csv.reader(fh))
            encabezado, datos = (filas[0] if filas else ()), filas[1:]
        if "Descripcion Operacion" not in encabezado:
            return True
        col = list(encabezado).index("Descripcion Operacion")
        return all(str(f[col] or "").upper().startswith(esperado) for f in datos if len(f) > col)
    except Exception:
        return True


async def bot_mis_retenciones_iva_simple(
    cuit_representante: str,
    clave_representante: str,
    cuit_representado: Optional[str],
    denominacion: str,
    desde: str,
    hasta: str,
    carga_minio: bool = True,
    proxy_request: bool = False,
) -> tuple[dict[str, Any], list[str], bool]:
    results: dict[str, Any] = {"archivos": []}
    error_logs: list[str] = []
    success = True

    if not cuit_representante:
        error_logs.append("El CUIT del representante es obligatorio")
        return results, error_logs, False
    if not (denominacion or "").strip():
        error_logs.append("La denominacion del representado es obligatoria")
        return results, error_logs, False
    if not desde or not hasta:
        error_logs.append("Las fechas desde y hasta son obligatorias")
        return results, error_logs, False
    date_error = validate_iva_simple_date_range(desde, hasta)
    if date_error:
        error_logs.append(date_error)
        return results, error_logs, False

    cuit_objetivo = cuit_representado or cuit_representante
    proxy = get_proxy_config() if proxy_request else None
    temp_dir = None
    base_root = Path("./descargas/mis_retenciones_iva_simple")
    if carga_minio:
        temp_dir = tempfile.TemporaryDirectory(prefix="mis_retenciones_iva_simple_")
        base_root = Path(temp_dir.name)

    stamp_dir = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_representado = safe_filename(cuit_objetivo or cuit_representante)
    output_dir = base_root / safe_representado / stamp_dir
    output_dir.mkdir(parents=True, exist_ok=True)

    async def subir_minio(path: Path, nombre_archivo: str) -> str:
        if not carga_minio:
            return ""
        try:
            return await subir_archivo_a_minio(
                minio_url_interna=minio_url_interna,
                minio_url_externa=minio_url_externa,
                minio_ssl=minio_ssl,
                minio_interno=minio_interno,
                minio_access_key=minio_access_key_iva_simple,
                minio_secret_key=minio_secret_key_iva_simple,
                minio_bucket=minio_bucket_iva_simple,
                archivo_a_subir=str(path),
                nombre_archivo_en_bucket=nombre_archivo,
            )
        except S3Error as exc:
            error_logs.append(f"Error al subir {path.name} a MinIO: {exc}")
        except Exception as exc:
            error_logs.append(f"Error al subir {path.name} a MinIO: {exc}")
        return ""

    try:
        async with limit_browser(), async_playwright() as playwright:
            browser_kwargs = {"headless": True, "proxy": proxy}
            executable_path = os.getenv("PLAYWRIGHT_EXECUTABLE_PATH")
            if executable_path:
                browser_kwargs["executable_path"] = executable_path
            browser = await playwright.chromium.launch(**browser_kwargs)
            context = await browser.new_context(accept_downloads=True)
            page = await context.new_page()

            async def cleanup_and_return(ok: bool) -> tuple[dict[str, Any], list[str], bool]:
                try:
                    await context.close()
                except Exception:
                    pass
                try:
                    await browser.close()
                except Exception:
                    pass
                return results, error_logs, ok

            login_result = await login_arca(
                page,
                cuit_representante,
                clave_representante,
                error_logs=error_logs,
            )
            if not login_result.success:
                return await cleanup_and_return(False)

            page1 = await open_arca_service(
                page,
                service_name="MIS RETENCIONES",
                service_matchers=(
                    re.compile("MIS RETENCIONES", re.I),
                    re.compile("MIS\\s+RETENCIONES|Mis\\s+Retenciones", re.I),
                ),
                error_logs=error_logs,
                ver_todos_error_message="No se pudo abrir 'Ver todos' en ARCA",
                service_unavailable_message="No se pudo abrir Mis Retenciones en ARCA",
            )
            if not page1:
                return await cleanup_and_return(False)

            await wait_for_ready(page1)
            await page1.goto("https://mirequa-web.arca.gob.ar/")
            if not await ensure_inicio(page1, cuit_objetivo):
                error_logs.append(
                    f"No se pudo seleccionar el representado {cuit_objetivo} en Mis Retenciones"
                )

            for operation in IVA_SIMPLE_OPERACIONES:
                for export_option, export_suffix in (
                    ("iva_simple", "iva_simple"),
                    ("xls", "xls"),
                ):
                    await page1.goto("https://mirequa-web.arca.gob.ar/")
                    await wait_for_ready(page1)
                    if not await ensure_inicio(page1, cuit_objetivo):
                        error_logs.append(
                            f"No se pudo confirmar el representado {cuit_objetivo} antes de consultar {operation}"
                        )
                        continue
                    if not await set_iva_simple_mode(page1):
                        error_logs.append("No se pudo seleccionar 'Exportar para IVA SIMPLE'")
                        continue
                    if not await select_iva_simple_operation(page1, operation):
                        error_logs.append(f"No se pudo seleccionar el tipo de operación {operation}")
                        continue
                    if not await ensure_iva_simple_date_range(page1, desde, hasta):
                        error_logs.append(
                            f"No se pudo cargar el rango de fechas {desde} - {hasta} para {operation}"
                        )
                        continue

                    nombre_archivo = build_retenciones_filename(
                        cuit_objetivo,
                        desde,
                        hasta,
                        IVA_SIMPLE_RESULT_LABEL,
                        operation,
                        denominacion,
                    )
                    nombre_archivo = nombre_archivo[:-4] + f" - {export_suffix}.csv"
                    csv_path = await download_csv(
                        page1,
                        "",
                        IVA_SIMPLE_RESULT_LABEL,
                        operation,
                        nombre_archivo,
                        output_dir,
                        error_logs,
                        required_result_marker=IVA_SIMPLE_RESULT_LABEL,
                        export_option=export_option,
                    )
                    if csv_path and not archivo_coincide_operacion(csv_path, operation):
                        error_logs.append(
                            f"El archivo {export_suffix} de {operation} no corresponde a esa operación; se descarta"
                        )
                        csv_path.unlink(missing_ok=True)
                        csv_path = None
                    if csv_path:
                        url_minio = await subir_minio(csv_path, csv_path.name) if carga_minio else ""
                        results["archivos"].append(
                            {
                                "impuesto_code": None,
                                "impuesto_label": IVA_SIMPLE_RESULT_LABEL,
                                "tipo": operation,
                                "tipo_operacion": operation,
                                "archivo": csv_path.name,
                                "formato": export_suffix,
                                "url_minio": url_minio or None,
                            }
                        )
                        if url_minio and csv_path.exists():
                            try:
                                csv_path.unlink()
                            except Exception as exc:
                                error_logs.append(
                                    f"Error al eliminar el archivo local {csv_path.name}: {exc}"
                                )

            await context.close()
            await browser.close()
    except Exception as exc:
        error_logs.append(f"Error en bot Mis Retenciones IVA SIMPLE: {exc}")
        success = False
    finally:
        if temp_dir:
            temp_dir.cleanup()

    if not results["archivos"] and not error_logs:
        error_logs.append("No se encontraron retenciones IVA SIMPLE para el periodo solicitado")
    return results, error_logs, success


async def main() -> None:
    from dotenv import load_dotenv
    load_dotenv()
    await bot_mis_retenciones_iva_simple(
        cuit_representante=os.getenv("CUIT_REPRESENTANTE"),
        clave_representante=os.getenv("CLAVE"),
        cuit_representado=os.getenv("CUIT_REPRESENTADO"),
        denominacion=os.getenv("DENOMINACION", "MI EMPRESA"),
        desde=os.getenv("DESDE"),
        hasta=os.getenv("HASTA"),
        carga_minio=False,
        proxy_request=False,
    )

if __name__ == "__main__":
    asyncio.run(main())
