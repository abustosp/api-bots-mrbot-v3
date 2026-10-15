import asyncio
import json
import logging
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from dotenv import load_dotenv
from playwright.async_api import (
    Browser,
    BrowserContext,
    Locator,
    Page,
    Playwright,
    TimeoutError as PlaywrightTimeoutError,
    async_playwright,
)

from shared.arca_login import launch_arca_browser_context, login_arca, open_arca_service, wait_page_ready

BOT_NAME = "carga_libro_iva"
TEMP_DIR = Path("./temp")
LOGS_ERROR_DIR = Path("./Logs_Error_Individuales")
INVALID_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')
PORTAL_IVA_URL_PATTERN = re.compile(r"(siapweb\.cloud\.afip\.gob\.ar/iva|liva\.afip\.gob\.ar)", re.I)
PORTAL_IVA_DIRECT_URL = "https://siapweb.cloud.afip.gob.ar/iva"
PORTAL_IVA_INIT_URL = "https://siapweb.cloud.afip.gob.ar/iva/#/init"
NO_PERIODOS_DISPONIBLES_RE = re.compile(
    r"No se encontraron per[ií]odos disponibles de presentaci[oó]n", re.I
)
TEMP_DIR.mkdir(exist_ok=True)
LOGS_ERROR_DIR.mkdir(exist_ok=True)

HEADLESS_MODE = os.getenv("CARGA_LIBRO_IVA_HEADLESS", "true").strip().lower() in ("true", "1", "yes")
LIVA_BASE = "https://liva.afip.gob.ar/liva/jsp"

# Proxy propia
USE_PROXY_PROPIA = os.getenv("CARGA_LIBRO_IVA_USE_PROXY_PROPIA", "false").strip().lower() in ("true", "1", "yes")
PROXY_PROPIA_IP = os.getenv("PROXY_PROPIA_IP", "").strip()
PROXY_PROPIA_PUERTO = os.getenv("PROXY_PROPIA_PUERTO", "").strip()
PROXY_PROPIA_USUARIO = os.getenv("PROXY_PROPIA_USUARIO", "").strip().strip('"')
PROXY_PROPIA_CLAVE = os.getenv("PROXY_PROPIA_CLAVE", "").strip().strip('"')


class ErrorNoReintentable(Exception):
    """Error que NO debe disparar reintentos del worker."""
    pass


class NoHayPeriodosDisponibles(ErrorNoReintentable):
    """El contribuyente no tiene períodos habilitados para presentar."""
    pass

VALORES_PROCESAR = {"si", "sí", "yes", "true", "1"}

# --- Timeouts ---
SHORT = 1_500
MEDIUM = 3_000
LONG = 5_000
CLICK_TO = 10_000
NAV_TO = 15_000
UPLOAD_TO = 60_000
MAX_RETRIES = 3


def _safe_step_name(step: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "_", step).strip("_")
    return safe or "step"


def _format_cuit(cuit: str) -> str:
    digits = re.sub(r"\D", "", cuit or "")
    if len(digits) == 11:
        return f"{digits[:2]}-{digits[2:10]}-{digits[10:]}"
    return cuit


def _sanitize(value: str) -> str:
    return INVALID_FILENAME_RE.sub("_", (value or "").strip())


def _cuit_folder_id(cuit: str) -> str:
    return re.sub(r"\D", "", cuit or "") or _sanitize(cuit) or "sin_cuit"


def _normalize_period(periodo: str) -> str:
    clean = (periodo or "").strip()
    match = re.fullmatch(r"(\d{2})/(\d{4})", clean)
    if match:
        return f"{match.group(2)}{match.group(1)}"
    match = re.fullmatch(r"(\d{4})(\d{2})", clean)
    if match:
        return clean
    match = re.fullmatch(r"(\d{2})-(\d{4})", clean)
    if match:
        return f"{match.group(2)}{match.group(1)}"
    return re.sub(r"\D", "", clean)


def _looks_like_blank_html(content: str) -> bool:
    compact = re.sub(r"\s+", "", content or "").lower()
    if not compact:
        return True
    if compact in {"<html><head></head><body></body></html>", "<html><head></head><body></body></html>"}:
        return True
    return len(compact) < 120


async def _is_blank_page(page: Page) -> bool:
    try:
        content = await page.content()
    except Exception:
        return True
    return _looks_like_blank_html(content)


def _portal_iva_init_url_with_key() -> str:
    token = format(int(datetime.now().timestamp() * 1000), "x")[-6:]
    return f"{PORTAL_IVA_INIT_URL}?_k={token}"


@dataclass
class ContribuyenteRow:
    nro: str
    cliente: str
    cuit_afip: str
    cuit_en_pagina: str
    clave: str
    cf: str
    tipo: str
    operaciones_ng_e: str
    periodo_afip: str
    ubicacion_txt: str
    txt_liv_c: str
    txt_liv_a: str
    txt_lic_c: str
    txt_lic_a: str
    directorio_iva_simple: str
    iva_simple: str

    @property
    def periodo_normalizado(self) -> str:
        return _normalize_period(self.periodo_afip)

    @property
    def cuit_formateado(self) -> str:
        return _format_cuit(self.cuit_en_pagina)

    @property
    def cuit_representante(self) -> str:
        """CUIT usado exclusivamente para iniciar sesión en ARCA."""
        return self.cuit_afip

    @property
    def cuit_representado(self) -> str:
        """CUIT del perfil sobre el que se ejecuta la carga."""
        return self.cuit_en_pagina

    @property
    def es_ng_e(self) -> bool:
        val = str(self.operaciones_ng_e or "").strip().upper()
        return val in ("NG/E", "NGE", "SI", "TRUE", "1")

    @property
    def prorrateo_tipo(self) -> Optional[str]:
        val = str(self.cf or "").strip().lower()
        if "global" in val:
            return "global"
        if "directa" in val or "comprobante" in val:
            return "directa"
        if "ambos" in val:
            return "ambos"
        return None


@dataclass
class ResultadoCarga:
    nro: str = ""
    cliente: str = ""
    cuit: str = ""
    periodo: str = ""
    venta_cbte: str = ""
    venta_alicuota: str = ""
    compra_cbte: str = ""
    compra_alicuota: str = ""
    cf_status: str = ""
    cf_restitucion: str = ""
    df_status: str = ""
    df_restitucion: str = ""
    estado_general: str = "PENDIENTE"
    error_msg: str = ""


# --- Column mapping ---
_EXCEL_COLUMNS = [
    ("Procesar", "Procesar"),
    ("Nro", "Nro"),
    ("Cliente", "Cliente"),
    ("CUIT AFIP", "CUIT AFIP"),
    ("CUIT en pagina", "CUIT en pagina"),
    ("CLAVE", "CLAVE"),
    ("CF", "CF"),
    ("Tipo", "Tipo"),
    ("Operaciones NG o E", "Operaciones NG o E"),
    ("Periodo AFIP", "Periodo AFIP"),
    ("Ubicacion TXT", "Ubicacion TXT"),
    ("TXT LIV C", "TXT LIV C"),
    ("TXT LIV A", "TXT LIV A"),
    ("TXT LIC C", "TXT LIC C"),
    ("TXT LIC A", "TXT LIC A"),
    ("Directorio IVA Simple", "Directorio IVA Simple"),
    ("IVA Simple", "IVA Simple"),
]


def _build_col_idx(headers: list[str]) -> dict[str, int]:
    norm_headers = [re.sub(r"[áéíóú]", lambda m: {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u"}[m.group(0)], h.lower().strip()) for h in headers]
    col_idx = {}
    for key, expected in _EXCEL_COLUMNS:
        norm_key = re.sub(r"[áéíóú]", lambda m: {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u"}[m.group(0)], key.lower().strip())
        try:
            idx = norm_headers.index(norm_key)
        except ValueError:
            for i, nh in enumerate(norm_headers):
                if norm_key in nh or nh in norm_key:
                    idx = i
                    break
            else:
                continue
        col_idx[key] = idx
    return col_idx


def _str_val(row, col_idx: dict[str, int], key: str) -> str:
    idx = col_idx.get(key)
    if idx is None:
        return ""
    return str(row[idx] or "").strip()


def leer_contribuyentes(excel_path: str) -> list[ContribuyenteRow]:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Listado"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    col_idx = _build_col_idx(headers)

    contribuyentes: list[ContribuyenteRow] = []
    for row in rows:
        procesar = _str_val(row, col_idx, "Procesar").lower()
        tipo = _str_val(row, col_idx, "Tipo").upper()
        if procesar not in VALORES_PROCESAR or tipo != "RI":
            continue

        c = ContribuyenteRow(
            nro=_str_val(row, col_idx, "Nro"),
            cliente=_str_val(row, col_idx, "Cliente"),
            cuit_afip=_str_val(row, col_idx, "CUIT AFIP"),
            cuit_en_pagina=_str_val(row, col_idx, "CUIT en pagina"),
            clave=_str_val(row, col_idx, "CLAVE"),
            cf=_str_val(row, col_idx, "CF"),
            tipo=_str_val(row, col_idx, "Tipo"),
            operaciones_ng_e=_str_val(row, col_idx, "Operaciones NG o E"),
            periodo_afip=_str_val(row, col_idx, "Periodo AFIP"),
            ubicacion_txt=_str_val(row, col_idx, "Ubicacion TXT"),
            txt_liv_c=_str_val(row, col_idx, "TXT LIV C"),
            txt_liv_a=_str_val(row, col_idx, "TXT LIV A"),
            txt_lic_c=_str_val(row, col_idx, "TXT LIC C"),
            txt_lic_a=_str_val(row, col_idx, "TXT LIC A"),
            directorio_iva_simple=_str_val(row, col_idx, "Directorio IVA Simple"),
            iva_simple=_str_val(row, col_idx, "IVA Simple"),
        )
        contribuyentes.append(c)

    wb.close()
    return contribuyentes


class DebugCapture:
    def __init__(self, prefix: str = BOT_NAME) -> None:
        run_id = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.run_dir = TEMP_DIR / f"{prefix}_{run_id}"
        self.logs_dir = self.run_dir / "logs"
        self.downloads_dir = self.run_dir / "descargas"
        self.screenshots_dir = self.run_dir / "screenshots"
        self.html_dir = self.run_dir / "html"
        self.har_dir = self.run_dir / "har"
        for path in (self.logs_dir, self.downloads_dir, self.screenshots_dir, self.html_dir, self.har_dir):
            path.mkdir(parents=True, exist_ok=True)

        self.step = 0
        self.logger = logging.getLogger(f"{prefix}.{run_id}")
        self.logger.setLevel(logging.INFO)
        self.logger.propagate = False

        formatter = logging.Formatter("%(asctime)s %(levelname)s %(message)s")
        fh = logging.FileHandler(self.logs_dir / f"{prefix}.log")
        fh.setFormatter(formatter)
        self.logger.addHandler(fh)
        sh = logging.StreamHandler()
        sh.setFormatter(formatter)
        self.logger.addHandler(sh)

    def log(self, message: str, *args) -> None:
        self.logger.info(message, *args)

    def log_error(self, message: str, *args) -> None:
        self.logger.error(message, *args)

    async def capture(self, page: Optional[Page], step: str) -> None:
        if page is None:
            return
        self.step += 1
        safe_step = _safe_step_name(step)
        pfx = f"{self.step:03d}_{safe_step}"
        try:
            await page.screenshot(path=str(self.screenshots_dir / f"{pfx}.png"), full_page=True)
        except Exception:
            pass
        try:
            (self.html_dir / f"{pfx}.html").write_text(await page.content(), encoding="utf-8")
        except Exception:
            pass


# --- Helper functions ---
async def try_click(locator: Locator, timeout: int = CLICK_TO) -> bool:
    try:
        await locator.first.click(timeout=timeout)
        return True
    except Exception:
        return False


async def click_first(locators: list[Locator], timeout: int = CLICK_TO) -> bool:
    for locator in locators:
        try:
            if await locator.count() == 0:
                continue
            if await try_click(locator, timeout=timeout):
                return True
        except Exception:
            continue
    return False


async def is_visible(locator: Locator, timeout: int = SHORT) -> bool:
    try:
        if await locator.count() == 0:
            return False
        await locator.first.wait_for(state="visible", timeout=timeout)
        return True
    except Exception:
        return False


async def click_modificar_si_corresponde(page: Page) -> bool:
    locators = [
        page.locator("#btnModificar"),
        page.get_by_role("button", name=re.compile(r"^MODIFICAR$", re.I)),
        page.get_by_role("link", name=re.compile(r"^MODIFICAR$", re.I)),
        page.get_by_text(re.compile(r"^MODIFICAR$", re.I), exact=True),
    ]
    for locator in locators:
        if await is_visible(locator, timeout=750) and await try_click(locator, timeout=CLICK_TO):
            await page.wait_for_timeout(SHORT)
            confirm_locators = [
                page.locator(".bootstrap-dialog .btn-primary").filter(
                    has_text=re.compile(r"^Continuar$", re.I)
                ),
                page.get_by_role("button", name=re.compile(r"^Continuar$", re.I)),
            ]
            for confirm_locator in confirm_locators:
                if await is_visible(confirm_locator, timeout=750):
                    await try_click(confirm_locator, timeout=CLICK_TO)
                    break
            return True
    return False


async def esperar_procesando(page: Page, debug: DebugCapture) -> None:
    try:
        procesando = page.get_by_text("Procesando", exact=False)
        if await is_visible(procesando, timeout=2_000):
            debug.log("procesando_detectado esperando...")
            try:
                await procesando.first.wait_for(state="hidden", timeout=UPLOAD_TO)
                debug.log("procesando_finalizo")
            except Exception:
                pass
            try:
                await page.wait_for_timeout(SHORT)
            except Exception:
                pass
    except Exception:
        pass


async def _no_hay_periodos_disponibles(page: Page) -> bool:
    """Detecta la baja o cualquier estado sin períodos habilitados."""
    try:
        mensaje = page.get_by_text(NO_PERIODOS_DISPONIBLES_RE, exact=False)
        if await is_visible(mensaje, timeout=1_000):
            return True
    except Exception:
        pass
    try:
        body_text = await page.locator("body").inner_text()
    except Exception:
        return False
    return bool(NO_PERIODOS_DISPONIBLES_RE.search(body_text))


async def execute_step(debug: DebugCapture, page: Optional[Page], step_name: str, action):
    debug.log("step_start=%s url=%s", step_name, getattr(page, "url", ""))
    try:
        result = await action()
        await debug.capture(page, step_name)
        debug.log("step_end=%s status=OK url=%s", step_name, getattr(page, "url", ""))
        return result
    except NoHayPeriodosDisponibles as exc:
        await debug.capture(page, f"{step_name}_sin_periodos")
        debug.log("step_end=%s status=SIN_PERIODOS url=%s mensaje=%s", step_name, getattr(page, "url", ""), exc)
        raise
    except Exception as exc:
        debug.log_error("step_end=%s status=ERROR url=%s error=%s", step_name, getattr(page, "url", ""), exc)
        await debug.capture(page, f"{step_name}_error")
        raise


async def ensure_session_active(page: Page) -> None:
    checks = [
        page.get_by_text("usuario no está logueado", exact=False),
        page.get_by_text("sesión ha expirado", exact=False),
        page.get_by_role("button", name=re.compile(r"IR A LA P[ÁA]GINA DE AFIP", re.I)),
    ]
    for locator in checks:
        if await is_visible(locator, timeout=500):
            raise RuntimeError("La sesión de Portal IVA expiró")


async def return_to_home(page: Page, debug: DebugCapture) -> None:
    await ensure_session_active(page)
    await esperar_procesando(page, debug)
    cur_url = getattr(page, "url", "")

    if "menuPresentacion" in cur_url:
        debug.log("navigation=already_at_home")
        return

    if "liva.afip.gob.ar" in cur_url:
        try:
            debug.log("navigation=return_via_goto")
            await page.goto(f"{LIVA_BASE}/menuPresentacion.do", wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
            await esperar_procesando(page, debug)
            if "menuPresentacion" in getattr(page, "url", ""):
                return
        except Exception:
            pass

    for locators in (
        [page.locator("a[href*='menuPresentacion.do']"), page.locator("#lnkInicio")],
        [page.get_by_role("link", name=re.compile(r"Inicio", re.I)), page.get_by_text("Inicio", exact=False)],
        [page.get_by_role("link", name=re.compile(r"Libro IVA - Borrador", re.I)), page.get_by_text("Libro IVA - Borrador", exact=False)],
        [page.get_by_role("button", name=re.compile(r"Volver", re.I)), page.get_by_text(re.compile(r"Volver", re.I)), page.locator("#btnVolver")],
    ):
        if await click_first(locators, timeout=LONG):
            debug.log("navigation=return_to_home")
            await page.wait_for_timeout(MEDIUM)
            await esperar_procesando(page, debug)
            return

    # Fallback: navegación directa
    try:
        debug.log("navigation=return_fallback_goto")
        await page.goto(f"{LIVA_BASE}/menuPresentacion.do", wait_until="domcontentloaded", timeout=NAV_TO)
        await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page, debug)
    except Exception:
        pass


async def esperar_arca_post_login(page: Page, debug: DebugCapture) -> None:
    """Espera la navegación asíncrona de ARCA antes de buscar servicios."""
    for _ in range(20):
        current_url = getattr(page, "url", "").lower()
        if "login" not in current_url:
            return
        if await is_visible(page.get_by_text(re.compile(r"Ver\s+todos", re.I)), timeout=500):
            return
        await page.wait_for_timeout(1_000)
    debug.log("arca_post_login_url=%s", getattr(page, "url", ""))
    raise RuntimeError("ARCA no completó la navegación posterior al inicio de sesión")


# --- Login & Portal ---
async def login_y_abrir_portal(
    playwright: Playwright,
    contrib: ContribuyenteRow,
    debug: DebugCapture,
) -> tuple[Browser, BrowserContext, Page]:
    har_path = debug.har_dir / f"{BOT_NAME}.har"
    context_kwargs = {"accept_downloads": True, "record_har_path": str(har_path), "record_har_content": "embed"}
    proxy_config = None
    if USE_PROXY_PROPIA:
        # Re-leer variables para que estén disponibles en cada invocación
        ip = os.getenv("PROXY_PROPIA_IP", "").strip()
        puerto = os.getenv("PROXY_PROPIA_PUERTO", "").strip()
        usuario = os.getenv("PROXY_PROPIA_USUARIO", "").strip().strip('"')
        clave = os.getenv("PROXY_PROPIA_CLAVE", "").strip().strip('"')
        proxy_config = {"server": f"http://{ip}:{puerto}",
                        "username": usuario, "password": clave}
        debug.log("proxy_propia_activa host=%s:%s user=%s", ip, puerto, usuario)
    browser, context, page = await launch_arca_browser_context(
        playwright, headless=HEADLESS_MODE,
        proxy=proxy_config,
        context_kwargs=context_kwargs,
    )
    page.set_default_timeout(30_000)

    async def _login():
        login_errors: list[str] = []
        try:
            debug.log("login_arca_cuit_inicio_sesion=%s", contrib.cuit_representante)
            result = await login_arca(
                page,
                contrib.cuit_representante,
                contrib.clave,
                error_logs=login_errors,
                fail_on_captcha=True,
            )
        except Exception as exc:
            joined = str(exc).lower()
            if "err_proxy" in joined or "proxy" in joined:
                ip = os.getenv("PROXY_PROPIA_IP", "").strip()
                puerto = os.getenv("PROXY_PROPIA_PUERTO", "").strip()
                raise RuntimeError(f"Error de conexión con proxy propia ({ip}:{puerto}): {exc}")
            raise
        for line in login_errors:
            debug.log("arca_login=%s", line)
        if result.captcha_detected:
            raise ErrorNoReintentable("Captcha detectado en ARCA. No se puede continuar automáticamente.")
        if result.password_change_required:
            raise ErrorNoReintentable("Se requiere cambiar la clave fiscal antes de continuar.")
        if await is_visible(
            page.get_by_text(re.compile(r"Clave\s+o\s+usuario\s+incorrecto", re.I)),
            timeout=750,
        ):
            raise ErrorNoReintentable(
                "Clave fiscal incorrecta para CUIT representante "
                f"{contrib.cuit_representante}"
            )
        if not result.success:
            joined = " | ".join(login_errors).lower()
            if any(k in joined for k in ("clave", "incorrect", "invalid", "no coincide", "inexistente", "usuario")):
                raise ErrorNoReintentable(
                    "Clave fiscal incorrecta para CUIT representante "
                    f"{contrib.cuit_representante}. Detalle: "
                    f"{login_errors[-1] if login_errors else 'credenciales inválidas'}"
                )
            if "err_proxy" in joined or "proxy" in joined:
                ip = os.getenv("PROXY_PROPIA_IP", "").strip()
                puerto = os.getenv("PROXY_PROPIA_PUERTO", "").strip()
                raise RuntimeError(f"Error de conexión con proxy propia ({ip}:{puerto}): {login_errors[-1] if login_errors else 'proxy no disponible'}")
            raise RuntimeError(login_errors[-1] if login_errors else "No se pudo iniciar sesión en ARCA")
        await esperar_arca_post_login(page, debug)

    await execute_step(debug, page, "login_arca", _login)

    async def _open_service():
        open_errors: list[str] = []
        portal_page = await open_arca_service(
            page, service_name="Portal IVA",
            service_matchers=(re.compile(r"PORTAL\s+IVA", re.I), "PORTAL IVA Presentación del"),
            error_logs=open_errors, click_ver_todos=True,
            service_unavailable_message="No se encuentra habilitado el servicio de Portal IVA",
        )
        for line in open_errors:
            debug.log("arca_service=%s", line)
        if portal_page is None:
            raise ErrorNoReintentable(open_errors[-1] if open_errors else "No se encuentra habilitado el servicio de Portal IVA")
        portal_page.set_default_timeout(30_000)

        try:
            await wait_page_ready(portal_page, timeout_ms=20_000)
            await portal_page.wait_for_url(PORTAL_IVA_URL_PATTERN, timeout=20_000)
        except PlaywrightTimeoutError:
            await portal_page.goto(PORTAL_IVA_DIRECT_URL, wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(portal_page, timeout_ms=20_000)
            await portal_page.wait_for_url(PORTAL_IVA_URL_PATTERN, timeout=20_000)

        if await _is_blank_page(portal_page):
            await portal_page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(portal_page, timeout_ms=20_000)

        await portal_page.wait_for_timeout(LONG)
        return portal_page

    portal_page = await execute_step(debug, page, "abrir_servicio_portal_iva", _open_service)
    return browser, context, portal_page


# --- Representado ---
async def seleccionar_representado(page: Page, contrib: ContribuyenteRow, debug: DebugCapture) -> None:
    formato = contrib.cuit_formateado
    target_digits = re.sub(r"\D", "", contrib.cuit_representado or "")
    login_digits = re.sub(r"\D", "", contrib.cuit_representante or "")
    active_profile_cuit = re.compile(rf"\[\s*CUIT\s+{re.escape(formato)}\s*\]", re.I)

    async def _is_target() -> bool:
        if await is_visible(
            page.locator("span.regular.text-uppercase").filter(has_text=active_profile_cuit),
            timeout=1_500,
        ):
            return True
        if await is_visible(page.get_by_text(active_profile_cuit, exact=False), timeout=1_500):
            return True
        if await is_visible(page.get_by_text(re.compile(rf"Representando a:.*{re.escape(formato)}", re.I)), timeout=1_500):
            return True
        if target_digits == login_digits and await is_visible(page.get_by_text(re.compile(re.escape(formato), re.I)), timeout=1_500):
            return True
        return False

    async def _select_from_panels() -> bool:
        """Busca en los paneles a.panel.panel-default y clickea el que tenga el CUIT buscado."""
        panel_locator = page.locator('a.panel.panel-default[title="Representar a..."]')
        count = await panel_locator.count()
        debug.log("paneles_representados_count=%d", count)
        for i in range(count):
            panel = panel_locator.nth(i)
            text = await panel.inner_text()
            if target_digits in text or formato in text:
                debug.log("clickeando_panel_representado_idx=%d cuit_encontrado", i)
                await panel.click(timeout=8_000)
                return True
        return False

    async def _select():
        await ensure_session_active(page)
        try:
            await wait_page_ready(page, timeout_ms=NAV_TO)
        except Exception:
            pass
        await page.wait_for_timeout(SHORT)

        for _ in range(2):
            if not await _is_blank_page(page):
                break
            await page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(page, timeout_ms=20_000)
            await page.wait_for_timeout(LONG)

        if await _is_blank_page(page):
            await page.goto(PORTAL_IVA_DIRECT_URL, wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(page, timeout_ms=20_000)
            await page.wait_for_timeout(LONG)

        if await _is_blank_page(page):
            raise RuntimeError("No se pudo cargar Portal IVA")

        # --- Caso 1: login CUIT == target CUIT ---
        if target_digits == login_digits:
            debug.log("login_cuit_equal_target cuit=%s", formato)
            await page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(page, timeout_ms=15_000)
            await page.wait_for_timeout(LONG)
            if await _is_blank_page(page):
                await page.goto(PORTAL_IVA_DIRECT_URL, wait_until="domcontentloaded", timeout=20_000)
                await wait_page_ready(page, timeout_ms=15_000)
                await page.wait_for_timeout(LONG)
            if await _is_target():
                debug.log("representado_ya_es_target_sin_panel cuit=%s", formato)
                return

        if await _is_target():
            debug.log("representado_already_selected cuit=%s", formato)
            return

        # --- Caso 2: abrir panel de cambio de relación ---
        debug.log("abriendo_panel_cambio_relacion")
        await click_first([
            page.get_by_title("cambio relación"),
            page.locator("[title*='cambio relación' i]"),
            page.get_by_role("link", name=re.compile(r"cambio.*relación|cambio.*relacion", re.I)),
            page.locator("a").filter(has_text=re.compile(r"cambio.*relación|cambio.*relacion", re.I)),
            page.locator("i.fa-users, i.fa-user-friends, i.fa-user").first,
        ], timeout=10_000)
        await page.wait_for_timeout(MEDIUM)

        # --- Click "cambio relación" para ir a la página de selección ---
        debug.log("navegando_a_pagina_cambio_relacion")
        await click_first([
            page.get_by_title("cambio relación"),
            page.locator("[title*='cambio relación' i]"),
            page.get_by_role("link", name=re.compile(r"cambio.*relación|cambio.*relacion", re.I)),
            page.locator("a").filter(has_text=re.compile(r"cambio.*relación|cambio.*relacion", re.I)),
            page.locator("i.fa-users, i.fa-user-friends, i.fa-user").first,
        ], timeout=10_000)
        await page.wait_for_timeout(LONG)

        # Esperar a que los paneles de representados se rendericen (SPA Angular)
        try:
            await page.locator('a.panel.panel-default[title="Representar a..."]').first.wait_for(
                state="visible", timeout=10_000)
            debug.log("paneles_representados_visibles")
        except Exception:
            debug.log("paneles_representados_no_visibles_continuando")

        # --- Seleccionar representado directamente desde los paneles ---
        cuit_limpio = re.sub(r"\D", "", formato)
        debug.log("seleccionando_representado cuit=%s cliente=%s", formato, contrib.cliente)

        selected = await _select_from_panels()

        if not selected:
            debug.log("paneles_no_coincidieron buscando_por_texto")
            selected = await click_first([
                page.locator('a, button, span, div, li, h3, small').filter(
                    has_text=re.compile(re.escape(cuit_limpio), re.I)),
                page.get_by_text(re.compile(re.escape(cuit_limpio), re.I)),
                page.locator('a, button, span, div, li').filter(
                    has_text=re.compile(re.escape(contrib.cliente), re.I)).first,
            ], timeout=8_000)

        if not selected:
            debug.log("representado_no_encontrado_por_cuit")
            try:
                body_text = await page.locator("body").inner_text()
                debug.log("body_text_muestra=%s", body_text[:600])
            except Exception:
                pass

        if selected:
            await page.wait_for_timeout(MEDIUM)

        # --- Refrescar Portal IVA para aplicar contexto ---
        debug.log("refrescando_portal_iva_para_aplicar_representado")
        await page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
        await wait_page_ready(page, timeout_ms=15_000)
        await page.wait_for_timeout(LONG)

        if await _is_blank_page(page):
            await page.goto(PORTAL_IVA_DIRECT_URL, wait_until="domcontentloaded", timeout=20_000)
            await wait_page_ready(page, timeout_ms=15_000)
            await page.wait_for_timeout(LONG)

        if await _is_target():
            debug.log("representado_seleccionado_exitosamente cuit=%s", formato)
            return

        raise RuntimeError(f"No se pudo seleccionar el representado {formato}")

    await execute_step(debug, page, "seleccionar_representado", _select)


# --- LIVA Navigation + Configuration ---
async def navegar_a_liva_y_configurar(page: Page, contrib: ContribuyenteRow, debug: DebugCapture) -> None:
    async def _navegar():
        await ensure_session_active(page)
        await esperar_procesando(page, debug)

        # 1) Click Ingresar/Nueva si es necesario
        if not await is_visible(page.locator("#periodo"), timeout=2_000):
            await click_first([
                page.get_by_role("button", name=re.compile(r"iva\.home\.btn\.nueva", re.I)),
                page.get_by_label("Sin texto (iva.home.btn.nueva"),
                page.get_by_role("button", name=re.compile(r"Ingresar", re.I)),
                page.get_by_text(re.compile(r"Nueva|Ingresar", re.I)),
            ], timeout=NAV_TO)
            await page.wait_for_timeout(MEDIUM)

        if await _no_hay_periodos_disponibles(page):
            raise NoHayPeriodosDisponibles("No hay periodos para presentar")

        # 2) Seleccionar período y validar
        if await page.locator("#periodo").count() > 0:
            periodo_normalizado = contrib.periodo_normalizado
            # El <option> puede mostrar su value (202606) o su texto (06/2026)
            # Comparamos varias formas para evitar falsos negativos
            opciones_texts = []
            opciones_values = []
            try:
                option_handles = await page.locator("#periodo option").element_handles()
                for oh in option_handles:
                    val = await oh.get_attribute("value") or ""
                    txt = await oh.inner_text()
                    opciones_values.append(val)
                    opciones_texts.append(txt)
                debug.log("periodos_disponibles texts=[%s] values=[%s] esperado=%s",
                          ",".join(opciones_texts), ",".join(opciones_values), periodo_normalizado)
            except Exception:
                pass
            # Ignorar placeholders para no continuar con un selector vacío.
            opciones_utiles = [
                (value, text) for value, text in zip(opciones_values, opciones_texts)
                if any(
                    re.fullmatch(r"\d{6}", _normalize_period(raw))
                    for raw in (value.strip(), text.strip())
                )
            ]
            if not opciones_utiles:
                raise NoHayPeriodosDisponibles("No hay periodos para presentar")

            # Aceptamos el período si está en values o si el texto coincide con period_afip.
            disponibles_normalizados = {
                _normalize_period(value) for value, text in opciones_utiles if value.strip()
            }
            disponibles_normalizados.update(
                _normalize_period(text) for value, text in opciones_utiles if text.strip()
            )
            disponible = periodo_normalizado in disponibles_normalizados
            if not disponible:
                raise ErrorNoReintentable(
                    f"No hay período {contrib.periodo_afip} disponible para presentar. "
                    f"Períodos habilitados: {','.join(text for value, text in opciones_utiles)}")
            await page.locator("#periodo").select_option(periodo_normalizado)
            debug.log("periodo_seleccionado=%s", periodo_normalizado)
            await page.wait_for_timeout(SHORT)
            await click_first([
                page.get_by_role("button", name=re.compile(r"Validar", re.I)),
                page.get_by_label("Sin texto (iva.btn.home.validar.periodo.alt)"),
            ], timeout=12_000)
            await page.wait_for_timeout(MEDIUM)
            await esperar_procesando(page, debug)

        if await click_modificar_si_corresponde(page):
            debug.log("iva_simple_modificar_detectado")
            await page.wait_for_timeout(MEDIUM)
            await esperar_procesando(page, debug)

        if "liva" in getattr(page, "url", ""):
            debug.log("ya_en_liva_tras_validar")
        else:
            # 3) Ir a Registración y declaración / LIVA
            debug.log("navegando a LIVA")
            await click_first([
                page.get_by_role("button", name=re.compile(r"Registraci[oó]n y declaraci[oó]n", re.I)),
                page.get_by_role("button", name=re.compile(r"Registraci[oó]n", re.I)),
                page.get_by_role("button", name=re.compile(r"iva\.btn\.home\.liva", re.I)),
                page.get_by_label("Sin texto (iva.btn.home.liva."),
                page.get_by_role("button", name=re.compile(r"Libro\s+IVA\s+Digital", re.I)),
                page.locator("button").filter(has_text=re.compile(r"Registraci[oó]n|Libro IVA|Declaraci[oó]n", re.I)),
            ], timeout=NAV_TO)
            await page.wait_for_timeout(LONG)
            await wait_page_ready(page, timeout_ms=20_000)
            await esperar_procesando(page, debug)

        # 4) Configurar opciones (NG/E, prorrateo, etc.)
        if "menuPresentacion" in getattr(page, "url", ""):
            debug.log("ya_en_menu_presentacion salteando configuracion")
            return

        if "vistaPreviaSimplificada" in getattr(page, "url", ""):
            debug.log("vista_simplificada clickeando Modificar")
            await click_modificar_si_corresponde(page)
            await page.wait_for_timeout(MEDIUM)
            await esperar_procesando(page, debug)
            if "menuPresentacion" in getattr(page, "url", ""):
                return

        if await is_visible(page.locator("#btnVistaPrevia"), timeout=3_000):
            debug.log("vista_previa_detectada no requiere config")
            return

        if contrib.es_ng_e:
            operaciones = page.locator("#operacionesNoGravadasExentas")
            if await operaciones.count() > 0:
                try:
                    await operaciones.check()
                except Exception:
                    await click_first([
                        page.locator("label").filter(has_text=re.compile(r"No Gravado", re.I)).locator("span"),
                        operaciones,
                    ], timeout=LONG)
            else:
                await click_first([
                    page.locator("label").filter(has_text=re.compile(r"No Gravado", re.I)).locator("span"),
                ], timeout=LONG)
            await page.wait_for_timeout(SHORT)
            prorrateo = contrib.prorrateo_tipo
            if prorrateo:
                if await page.locator("#tipoProrrateo").count() > 0:
                    try:
                        await page.locator("#tipoProrrateo").select_option(
                            "3" if prorrateo == "ambos" else ("2" if prorrateo == "global" else "1"))
                    except Exception:
                        pass
                else:
                    await click_first([page.get_by_role("button", name=re.compile(r"prorrateo", re.I))], timeout=LONG)
                    await page.wait_for_timeout(SHORT)
                    opts = {
                        "global": ["Con prorrateo global"],
                        "directa": ["Con prorrateo por asignación directa"],
                        "ambos": ["Con prorrateo por asignación directa y global"],
                    }
                    await click_first([
                        page.get_by_role("option", name=re.compile(rf"{opts[prorrateo][0]}", re.I)),
                    ], timeout=LONG)
        else:
            ninguna = page.locator("#ningunaOperacion")
            if await ninguna.count() > 0:
                try:
                    await ninguna.check()
                except Exception:
                    await click_first([
                        page.locator("label").filter(has_text=re.compile(r"No realizo ninguna", re.I)).locator("span"),
                        ninguna,
                    ], timeout=LONG)
            else:
                await click_first([
                    page.locator("label").filter(has_text=re.compile(r"No realizo ninguna", re.I)).locator("span"),
                ], timeout=LONG)

        await page.wait_for_timeout(SHORT)
        guardar = page.locator("#btnGuardar").first
        if await guardar.count() > 0:
            if not await guardar.is_enabled():
                raise RuntimeError("No se pudo habilitar Continuar en la configuración de IVA")
            await guardar.click(timeout=CLICK_TO)
            await page.wait_for_timeout(MEDIUM)

    await execute_step(debug, page, "navegar_a_liva_y_configurar", _navegar)


# --- Libros: abrir, eliminar, importar ---
async def abrir_libro(page: Page, book: str, debug: DebugCapture) -> None:
    href = "verVentas.do" if book == "ventas" else "verCompras.do"
    label = "Libro Ventas" if book == "ventas" else "Libro Compras"

    async def _abrir():
        await ensure_session_active(page)
        await esperar_procesando(page, debug)
        clicked = await click_first([
            page.locator(f"a[href*='{href}']"),
            page.get_by_role("link", name=re.compile(rf"{label}\s+Registr", re.I)),
            page.get_by_role("link", name=re.compile(label, re.I)),
            page.locator(f"#{'btnLibroVentas' if book == 'ventas' else 'btnLibroCompras'}"),
        ], timeout=NAV_TO)
        if not clicked:
            raise RuntimeError(f"No se pudo abrir {label}")
        await page.wait_for_timeout(MEDIUM)

    await execute_step(debug, page, f"abrir_{book}", _abrir)


async def eliminar_todo(page: Page, debug: DebugCapture) -> bool:
    async def _eliminar():
        await ensure_session_active(page)
        await page.wait_for_timeout(MEDIUM)
        deleted = await click_first([
            page.get_by_role("button", name=re.compile(r"Eliminar\s+Todos?\.\.\.?", re.I)),
            page.locator("#btnEliminarTodo"),
        ], timeout=8_000)
        if not deleted:
            return False
        await page.wait_for_timeout(SHORT)
        confirmed = await click_first([
            page.locator(".bootstrap-dialog .btn-danger").filter(has_text=re.compile(r"Eliminar", re.I)),
            page.locator(".modal.in .btn-danger"),
            page.get_by_role("button", name=re.compile(r"Eliminar", re.I)),
        ], timeout=LONG)
        if not confirmed:
            await click_first([page.locator(".bootstrap-dialog .close")], timeout=MEDIUM)
            return False
        await page.wait_for_timeout(MEDIUM)
        return True

    return await execute_step(debug, page, "eliminar_todo", _eliminar)


async def importar_txts(page: Page, book: str, cbte_path: str, alicuota_path: str, debug: DebugCapture) -> None:
    comp_label = "Archivo Ventas *" if book == "ventas" else "Archivo Compras *"
    alic_label = "Archivo IVA Ventas *" if book == "ventas" else "Archivo IVA Compras *"

    async def _importar():
        await ensure_session_active(page)
        opened = await click_first([
            page.locator("#btnDropdownImportar"),
            page.get_by_role("button", name=re.compile(r"Importar", re.I)),
        ], timeout=CLICK_TO)
        if not opened:
            raise RuntimeError(f"No se encontró Importar para {book}")
        await page.wait_for_timeout(SHORT)

        clicked = await click_first([
            page.get_by_role("link", name=re.compile(r"Importar\s+(?:desde\s+)?Archivos?", re.I)),
            page.locator("#lnkImportarArchivo"),
        ], timeout=12_000)
        if not clicked:
            raise RuntimeError(f"No se encontró 'Importar Archivos' para {book}")
        await page.wait_for_timeout(MEDIUM)

        await page.get_by_label(comp_label).set_input_files(cbte_path)
        debug.log("upload_cbte book=%s", book)
        await page.get_by_label(alic_label).set_input_files(alicuota_path)
        debug.log("upload_alicuota book=%s", book)

        await page.wait_for_timeout(SHORT)
        await page.locator("#btnImportarArchivosImportar").evaluate("el => el.click()")
        await page.wait_for_timeout(MEDIUM)

        await esperar_procesando(page, debug)
        try:
            await page.get_by_text("Procesada", exact=False).first.wait_for(state="visible", timeout=UPLOAD_TO)
            debug.log("procesado_detectado tras importacion %s", book)
            await click_first([
                page.get_by_text("Procesada", exact=False),
                page.locator(".bootstrap-dialog .close"),
            ], timeout=LONG)
        except Exception:
            pass
        await page.wait_for_timeout(SHORT)
        await click_first([page.locator("#btnTareasCerrar")], timeout=LONG)

    await execute_step(debug, page, f"importar_txts_{book}", _importar)


# --- Helper: Crédito Fiscal Prorrateo ---
async def _handle_credito_fiscal_modal(page: Page, debug: DebugCapture) -> bool:
    """Si aparece el modal de crédito fiscal con el campo #credFiscProrrateo,
    extrae el valor del texto visible, lo normaliza y lo ingresa."""
    input_field = page.locator("#credFiscProrrateo")
    if await input_field.count() == 0:
        return False
    debug.log("credito_fiscal_modal_detectado")
    # Buscar el valor en todo el texto visible de la página
    try:
        body_text = await page.locator("body").inner_text()
    except Exception:
        body_text = ""
    # Buscar patrones como "$ 1.234,56" o "1.234,56" o "Crédito Fiscal: 1234,56"
    numeros = re.findall(r'(?:\$?\s*)([\d]{1,3}(?:\.\d{3})*(?:,\d{2})?)', body_text)
    if not numeros:
        numeros = re.findall(r'(\d{1,3}(?:\.\d{3})*(?:,\d{2})?)', body_text)
    if numeros:
        raw = numeros[-1]
        debug.log("credito_fiscal_valor_raw=%s", raw)
        # Limpiar: sacar puntos (separador miles), coma → punto
        limpio = raw.replace(".", "").replace(",", ".")
        debug.log("credito_fiscal_valor_limpio=%s", limpio)
        await input_field.first.fill(limpio)
        debug.log("credito_fiscal_valor_ingresado")
        await page.wait_for_timeout(SHORT)
        return True
    debug.log("credito_fiscal_valor_no_encontrado")
    return False


async def _ir_a_editar_credito_fiscal_en_compras(page: Page, debug: DebugCapture) -> bool:
    """Navega a la sección 'Crédito Fiscal Computable Global' en libro compras
    y si encuentra el campo #credFiscProrrateo, intenta cargar el valor detectado."""
    await return_to_home(page, debug)
    await esperar_procesando(page, debug)
    await abrir_libro(page, "compras", debug)
    await esperar_procesando(page, debug)

    # Buscar link "Crédito Fiscal Computable Global" en verCompras.do
    clicked = await click_first([
        page.locator("a, button, span, div").filter(has_text=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
        page.get_by_role("link", name=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
        page.get_by_text(re.compile(r"Crédito\s+Fiscal\s+Computable", re.I), exact=False),
    ], timeout=8_000)
    if not clicked:
        debug.log("no_en_verCompras_buscando_en_menuInfoAdicional")
        # Fallback: ir a menuInfoAdicional y buscar allí
        await return_to_home(page, debug)
        await esperar_procesando(page, debug)
        try:
            await page.locator("#btnAjustes").first.click(timeout=CLICK_TO)
            await page.wait_for_timeout(MEDIUM)
        except Exception:
            pass
        try:
            if await page.locator("a[href*='menuInfoAdicional']").count() > 0:
                await page.locator("a[href*='menuInfoAdicional']").first.click(timeout=CLICK_TO)
                await page.wait_for_timeout(MEDIUM)
            else:
                await page.goto(f"{LIVA_BASE}/menuInfoAdicional.jsp", wait_until="domcontentloaded", timeout=NAV_TO)
                await wait_page_ready(page, timeout_ms=10_000)
        except Exception:
            await page.goto(f"{LIVA_BASE}/menuInfoAdicional.jsp", wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page, debug)
        clicked = await click_first([
            page.locator("a, button, span, div").filter(has_text=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
            page.get_by_role("link", name=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
            page.get_by_text(re.compile(r"Crédito\s+Fiscal", re.I), exact=False),
            page.locator("a.btn-primary").filter(has_text=re.compile(r"Crédito\s+Fiscal", re.I)),
            page.locator("a[href*='creditoFiscal']"),
        ], timeout=8_000)
    if not clicked:
        debug.log("no_se_encontro_credito_fiscal_computable_global")
        return False
    debug.log("navegado_a_credito_fiscal_computable_global")
    await page.wait_for_timeout(MEDIUM)
    await esperar_procesando(page, debug)
    manejado = await _handle_credito_fiscal_modal(page, debug)
    if not manejado:
        debug.log("no_hubo_modal_credito_fiscal_en_compras")
    # Guardar
    await click_first([
        page.locator("button[type='submit']"),
        page.get_by_role("button", name=re.compile(r"Guardar", re.I)),
    ], timeout=8_000)
    await page.wait_for_timeout(MEDIUM)
    return manejado


# --- Apertura CSV ---
async def cargar_apertura_seccion(page: Page, seccion: str, csv_path: str, debug: DebugCapture) -> str:
    if not os.path.isfile(csv_path):
        debug.log("csv_no_encontrado seccion=%s", seccion)
        return "Sin archivo"
    if os.path.getsize(csv_path) == 0:
        debug.log("csv_vacio seccion=%s", seccion)
        return "Sin movimientos"
    try:
        with open(csv_path, encoding="utf-8-sig") as f:
            non_empty_lines = [line for line in f if line.strip()]
        if len(non_empty_lines) <= 1:
            debug.log("csv_solo_header seccion=%s lineas=%d", seccion, len(non_empty_lines))
            return "Sin movimientos"
    except Exception:
        pass

    links_map = {
        "CF": re.compile(r"Crédito\s+Fiscal(?!\s+Comp)", re.I),
        "CF Restitucion": re.compile(r"Restituci[oó]n.*Crédito\s+Fiscal", re.I),
        "DF": re.compile(r"Débito\s+Fiscal(?!\s+Comp)", re.I),
        "DF Restitucion": re.compile(r"Restituci[oó]n.*Débito\s+Fiscal", re.I),
    }
    link_pattern = links_map.get(seccion)

    async def _safe_goto(page: Page, url: str) -> bool:
        """Navigate and wait; return False if page is gone."""
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
            return True
        except Exception:
            return False

    async def _navegar_a_menu_info() -> bool:
        """Navega a menuInfoAdicional, recuperándose de página en blanco."""
        cur = getattr(page, "url", "")
        if "menuInfoAdicional" in cur:
            return True

        # Si la página está en blanco (Portal IVA sin SPA), ir directo
        if await _is_blank_page(page):
            debug.log("blank_page_detected goto=menuInfoAdicional")
            return await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp")

        await esperar_procesando(page, debug)

        # Intentar navegación por clicks
        for attempt in range(2):
            try:
                if await page.locator("#btnAjustes").count() > 0:
                    await page.locator("#btnAjustes").first.click(timeout=CLICK_TO)
                    await page.wait_for_timeout(MEDIUM)
                if await page.locator("a[href*='menuInfoAdicional']").count() > 0:
                    await page.locator("a[href*='menuInfoAdicional']").first.click(timeout=CLICK_TO)
                    await page.wait_for_timeout(MEDIUM)
                    if "menuInfoAdicional" in getattr(page, "url", ""):
                        return True
                    continue
                if await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp"):
                    return True
            except Exception:
                if attempt == 1:
                    return await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp")
                continue

        # Fallback: goto directo
        if "menuInfoAdicional" not in getattr(page, "url", ""):
            debug.log("navegacion_fallback_menuInfoAdicional")
            return await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp")
        return True

    async def _cargar():
        if not await _navegar_a_menu_info():
            raise RuntimeError(f"No se pudo navegar a menuInfoAdicional para {seccion}")
        await esperar_procesando(page, debug)

        # Ir a la sección específica (ej: CF, DF, etc.)
        href_map = {
            "CF": "d=c", "CF Restitucion": "d=rc",
            "DF": "d=d", "DF Restitucion": "d=rd",
        }
        clicked = await click_first([
            page.locator(f".form-group a.btn-primary[href*='{href_map[seccion]}']"),
            page.locator("a.btn-primary").filter(has_text=link_pattern),
            page.get_by_role("link", name=link_pattern),
        ], timeout=NAV_TO)
        if not clicked:
            raise RuntimeError(f"No se encontró link para {seccion}")
        await page.wait_for_timeout(MEDIUM)
        await esperar_procesando(page, debug)

        # Click Importar
        btn_importar = page.locator("#btnImportarArchivo")
        if await btn_importar.count() > 0:
            try:
                await btn_importar.first.click(timeout=CLICK_TO)
            except Exception:
                await btn_importar.first.evaluate("el => el.click()")
        else:
            await click_first([
                page.locator("button.btn-link").filter(has_text=re.compile(r"Importar", re.I)),
                page.get_by_role("button", name=re.compile(r"Importar.*CSV|Importar.*Archivo", re.I)),
            ], timeout=CLICK_TO)
        await page.wait_for_timeout(MEDIUM)

        # Upload CSV
        file_input = page.locator("input[type='file']")
        if await file_input.count() > 0:
            await file_input.first.set_input_files(csv_path)
            debug.log("upload_csv seccion=%s", seccion)
        await page.wait_for_timeout(SHORT)

        # Importar archivo
        importar_btn = page.locator("#btnImportarArchivoImportar")
        if await importar_btn.count() > 0:
            try:
                await importar_btn.first.click(timeout=CLICK_TO)
            except Exception:
                try:
                    await importar_btn.evaluate("el => el.click()")
                except Exception:
                    pass
        await page.wait_for_timeout(MEDIUM)

        # Esperar procesamiento (con tolerancia a cierre de página)
        try:
            await esperar_procesando(page, debug)
        except Exception:
            pass
        try:
            await page.wait_for_timeout(2_000)
        except Exception:
            pass

        # Si es CF, manejar posible modal de crédito fiscal prorrateo
        if seccion == "CF":
            try:
                await _handle_credito_fiscal_modal(page, debug)
            except Exception as exc:
                debug.log("credito_fiscal_modal_warning=%s", exc)

        # Guardar
        try:
            await click_first([
                page.locator("button[type='submit']"),
                page.get_by_role("button", name=re.compile(r"Guardar", re.I)),
            ], timeout=8_000)
            await page.wait_for_timeout(MEDIUM)
        except Exception:
            pass
        try:
            await esperar_procesando(page, debug)
        except Exception:
            pass

        # Volver
        try:
            await click_first([
                page.locator("a[href*='menuInfoAdicional']"),
                page.locator("#btnVolver"),
                page.get_by_text(re.compile(r"Volver", re.I)),
                page.locator("a[href*='menuPresentacion']"),
            ], timeout=8_000)
            await page.wait_for_timeout(MEDIUM)
        except Exception:
            pass

    await execute_step(debug, page, f"cargar_apertura_{_safe_step_name(seccion)}", _cargar)
    return "OK"


# --- Procesamiento por contribuyente ---
async def procesar_contribuyente(playwright: Playwright, contrib: ContribuyenteRow) -> ResultadoCarga:
    res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado, periodo=contrib.periodo_afip)
    cuit_inicio_sesion_folder = _cuit_folder_id(contrib.cuit_representante)
    cuit_representado_folder = _cuit_folder_id(contrib.cuit_representado)
    debug_prefix = f"{BOT_NAME}_{cuit_inicio_sesion_folder}_{cuit_representado_folder}_{contrib.nro}"
    debug = DebugCapture(prefix=debug_prefix)
    browser: Optional[Browser] = None
    context: Optional[BrowserContext] = None
    portal_page: Optional[Page] = None

    try:
        debug.log(
            "inicio nro=%s cliente=%s cuit_inicio_sesion=%s "
            "cuit_representado=%s periodo=%s",
            contrib.nro,
            contrib.cliente,
            contrib.cuit_representante,
            contrib.cuit_representado,
            contrib.periodo_afip,
        )

        browser, context, portal_page = await login_y_abrir_portal(playwright, contrib, debug)
        await seleccionar_representado(portal_page, contrib, debug)
        await navegar_a_liva_y_configurar(portal_page, contrib, debug)

        if "liva" not in getattr(portal_page, "url", ""):
            raise RuntimeError(f"No se pudo acceder al sistema Libro IVA. URL: {getattr(portal_page, 'url', '')}")

        # --- TXTs Ventas ---
        txt_liv_c = os.path.join(contrib.ubicacion_txt, contrib.txt_liv_c) if contrib.ubicacion_txt and contrib.txt_liv_c else ""
        txt_liv_a = os.path.join(contrib.ubicacion_txt, contrib.txt_liv_a) if contrib.ubicacion_txt and contrib.txt_liv_a else ""

        if contrib.txt_liv_c and contrib.txt_liv_a and os.path.isfile(txt_liv_c) and os.path.isfile(txt_liv_a):
            if os.path.getsize(txt_liv_c) == 0 and os.path.getsize(txt_liv_a) == 0:
                res.venta_cbte = "Sin movimientos"
                res.venta_alicuota = "Sin movimientos"
            else:
                try:
                    await abrir_libro(portal_page, "ventas", debug)
                    await eliminar_todo(portal_page, debug)
                    await importar_txts(portal_page, "ventas", txt_liv_c, txt_liv_a, debug)
                    res.venta_cbte = "OK"
                    res.venta_alicuota = "OK"
                except Exception as e:
                    debug.log_error("ventas_error=%s", e)
                    res.venta_cbte = "ERROR"
                    res.venta_alicuota = "ERROR"
        else:
            res.venta_cbte = "Sin archivo"
            res.venta_alicuota = "Sin archivo"

        await return_to_home(portal_page, debug)

        # --- TXTs Compras ---
        txt_lic_c = os.path.join(contrib.ubicacion_txt, contrib.txt_lic_c) if contrib.ubicacion_txt and contrib.txt_lic_c else ""
        txt_lic_a = os.path.join(contrib.ubicacion_txt, contrib.txt_lic_a) if contrib.ubicacion_txt and contrib.txt_lic_a else ""

        if contrib.txt_lic_c and contrib.txt_lic_a and os.path.isfile(txt_lic_c) and os.path.isfile(txt_lic_a):
            if os.path.getsize(txt_lic_c) == 0 and os.path.getsize(txt_lic_a) == 0:
                res.compra_cbte = "Sin movimientos"
                res.compra_alicuota = "Sin movimientos"
            else:
                try:
                    await abrir_libro(portal_page, "compras", debug)
                    await eliminar_todo(portal_page, debug)
                    await importar_txts(portal_page, "compras", txt_lic_c, txt_lic_a, debug)
                    res.compra_cbte = "OK"
                    res.compra_alicuota = "OK"
                except Exception as e:
                    debug.log_error("compras_error=%s", e)
                    res.compra_cbte = "ERROR"
                    res.compra_alicuota = "ERROR"
        else:
            res.compra_cbte = "Sin archivo"
            res.compra_alicuota = "Sin archivo"

        await return_to_home(portal_page, debug)

        # --- Aperturas CSV ---
        await esperar_procesando(portal_page, debug)
        for seccion_key, res_attr in [
            ("CF", "cf_status"), ("CF Restitucion", "cf_restitucion"),
            ("DF", "df_status"), ("DF Restitucion", "df_restitucion"),
        ]:
            csv_filename = contrib.iva_simple.replace("(XXXX)", seccion_key)
            csv_path = os.path.join(contrib.directorio_iva_simple, csv_filename) if contrib.directorio_iva_simple else ""
            try:
                status = await cargar_apertura_seccion(portal_page, seccion_key, csv_path, debug)
                setattr(res, res_attr, status)
                if seccion_key == "CF" and status == "OK":
                    try:
                        await _ir_a_editar_credito_fiscal_en_compras(portal_page, debug)
                    except Exception as exc2:
                        debug.log("credito_fiscal_compras_warning=%s", exc2)
            except Exception as e:
                debug.log_error("apertura_error seccion=%s error=%s", seccion_key, e)
                setattr(res, res_attr, "ERROR")
                res.error_msg = str(e)
            # Siempre retornar a home entre secciones
            try:
                await return_to_home(portal_page, debug)
            except Exception:
                pass

        # --- Estado general ---
        estados = [res.venta_cbte, res.venta_alicuota, res.compra_cbte, res.compra_alicuota,
                   res.cf_status, res.cf_restitucion, res.df_status, res.df_restitucion]
        if any(e == "ERROR" for e in estados if e):
            res.estado_general = "ERROR"
        elif any(e == "Sin movimientos" for e in estados if e):
            res.estado_general = "PARCIAL (Sin movimientos)"
        elif any(e == "Sin archivo" for e in estados if e):
            res.estado_general = "PARCIAL (Faltan archivos)"
        else:
            res.estado_general = "COMPLETO"

        debug.log("resultado nro=%s estado=%s", contrib.nro, res.estado_general)

    except NoHayPeriodosDisponibles as e:
        debug.log("sin_periodos nro=%s mensaje=%s", contrib.nro, e)
        res.estado_general = "SIN PERIODOS"
        res.error_msg = str(e)
    except Exception as e:
        debug.log_error("error nro=%s error=%s", contrib.nro, e)
        res.estado_general = "ERROR"
        res.error_msg = str(e)
        error_log = {
            "nro": contrib.nro,
            "cliente": contrib.cliente,
            "cuit_inicio_sesion": contrib.cuit_representante,
            "cuit_representado": contrib.cuit_representado,
            "periodo": contrib.periodo_afip,
            "error": str(e),
            "fecha_ejecucion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }
        with open(
            LOGS_ERROR_DIR
            / f"log_{cuit_inicio_sesion_folder}_{cuit_representado_folder}_"
            f"{contrib.nro}_{contrib.periodo_normalizado}.json",
            "w",
        ) as f:
            json.dump(error_log, f, ensure_ascii=False, indent=4)

    finally:
        if context:
            try:
                await context.close()
            except Exception:
                pass
        if browser:
            try:
                await browser.close()
            except Exception:
                pass
        print(f"[{res.estado_general}] {contrib.cliente} ({res.cuit}) - {contrib.periodo_afip}")
        if res.error_msg:
            label = "Detalle" if res.estado_general == "SIN PERIODOS" else "Error"
            print(f"   {label}: {res.error_msg[:200]}")

    return res


# --- Worker con retry ---
async def worker(semaphore: asyncio.Semaphore, contrib: ContribuyenteRow,
                  resultados: list[ResultadoCarga], timeout: int = 300) -> None:
    async with semaphore:
        for intento in range(1, MAX_RETRIES + 1):
            print(f"[INICIANDO intento {intento}/{MAX_RETRIES}] {contrib.cliente} ({contrib.cuit_formateado}) - {contrib.periodo_afip}")
            try:
                async with async_playwright() as playwright:
                    res = await asyncio.wait_for(procesar_contribuyente(playwright, contrib), timeout=timeout)
                if res.estado_general == "SIN PERIODOS":
                    resultados.append(res)
                    return
                if res.estado_general == "ERROR" and intento < MAX_RETRIES:
                    # Si el error_msg indica "no reintentable", no reintentar
                    if any(k in (res.error_msg or "").lower() for k in (
                        "captcha detectado", "cambiar la clave fiscal", "clave fiscal incorrecta",
                        "no se encuentra habilitado el servicio", "no hay período",
                        "no hay periodo", "sin períodos", "no se encontraron períodos"
                    )):
                        print(f"[NO-REINTENTABLE] {contrib.cliente} - {res.error_msg[:120]}")
                        resultados.append(res)
                        return
                    print(f"[REINTENTO {intento}] {contrib.cliente} - error: {res.error_msg[:100]}")
                    await asyncio.sleep(3)
                    continue
                resultados.append(res)
                return
            except asyncio.TimeoutError:
                if intento < MAX_RETRIES:
                    print(f"[TIMEOUT intento {intento}] {contrib.cliente} - reintentando...")
                    await asyncio.sleep(3)
                    continue
                print(f"[TIMEOUT] {contrib.cliente}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="TIMEOUT", error_msg=f"Timeout {timeout}s tras {MAX_RETRIES} intentos")
                resultados.append(res)
                return
            except NoHayPeriodosDisponibles as e:
                print(f"[SIN PERIODOS] {contrib.cliente} - {e}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="SIN PERIODOS", error_msg=str(e))
                resultados.append(res)
                return
            except ErrorNoReintentable as e:
                print(f"[NO-REINTENTABLE] {contrib.cliente} - {e}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="ERROR", error_msg=str(e))
                resultados.append(res)
                return
            except Exception as e:
                if intento < MAX_RETRIES:
                    print(f"[REINTENTO {intento}] {contrib.cliente} - {e}")
                    await asyncio.sleep(3)
                    continue
                print(f"[EXCEPCION] {contrib.cliente} - {e}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="ERROR", error_msg=str(e))
                resultados.append(res)
                return


# --- Informe ---
def generar_informe(resultados: list[ResultadoCarga], output_path: str) -> None:
    import pandas as pd
    headers = ["#", "Contribuyente", "CUIT", "Período",
               "Venta CBTE", "Venta Alic.", "Compra CBTE", "Compra Alic.",
               "Apertura CF", "Apertura CF Rest.", "Apertura DF", "Apertura DF Rest.", "Estado"]
    rows = []
    errores = []
    for i, r in enumerate(resultados, 1):
        rows.append([i, r.cliente, r.cuit, r.periodo,
                     r.venta_cbte, r.venta_alicuota, r.compra_cbte, r.compra_alicuota,
                     r.cf_status, r.cf_restitucion, r.df_status, r.df_restitucion, r.estado_general])
        if r.estado_general == "ERROR" and r.error_msg:
            errores.append([r.cliente, r.cuit, "", r.error_msg, ""])
    df_resumen = pd.DataFrame(rows, columns=headers)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="Resumen", index=False)
        if errores:
            pd.DataFrame(errores, columns=["Contribuyente", "CUIT", "Paso", "Error", "Fecha"]).to_excel(
                writer, sheet_name="Detalle Errores", index=False)
    print(f"\nInforme guardado en: {output_path}")


# --- Main ---
async def main_batch(excel_path: str, max_workers: int = 4, timeout: int = 300,
                     output_informe: str = "informe_carga_libro_iva.xlsx") -> list[ResultadoCarga]:
    contribuyentes = leer_contribuyentes(excel_path)
    if not contribuyentes:
        print("No se encontraron contribuyentes para procesar.")
        return []

    print(f"Contribuyentes a procesar: {len(contribuyentes)}")
    for c in contribuyentes:
        print(f"  - [{c.nro}] {c.cliente} ({_format_cuit(c.cuit_en_pagina)}) periodo={c.periodo_afip} NG/E={c.es_ng_e} CF={c.prorrateo_tipo}")

    semaphore = asyncio.Semaphore(max_workers)
    resultados: list[ResultadoCarga] = []
    tasks = [asyncio.create_task(worker(semaphore, c, resultados, timeout)) for c in contribuyentes]
    await asyncio.gather(*tasks)
    resultados.sort(key=lambda r: int(r.nro) if r.nro.isdigit() else 999999)
    generar_informe(resultados, output_informe)

    completos = sum(1 for r in resultados if r.estado_general == "COMPLETO")
    parciales = sum(1 for r in resultados if r.estado_general.startswith("PARCIAL"))
    sin_periodos = sum(1 for r in resultados if r.estado_general == "SIN PERIODOS")
    errores = sum(1 for r in resultados if r.estado_general in ("ERROR", "TIMEOUT"))
    print(f"\nResumen: {completos} completos, {parciales} parciales, {sin_periodos} sin períodos, {errores} errores/timeouts")
    return resultados


if __name__ == "__main__":
    load_dotenv()
    EXCEL_PATH = os.getenv("CARGA_LIBRO_IVA_EXCEL", "Lista libros SOS.xlsx").strip()
    MAX_WORKERS = int(os.getenv("CARGA_LIBRO_IVA_MAX_WORKERS", "1"))
    TIMEOUT = int(os.getenv("CARGA_LIBRO_IVA_TIMEOUT", "600"))
    OUTPUT = os.getenv("CARGA_LIBRO_IVA_OUTPUT", "informe_carga_libro_iva.xlsx").strip()

    print("=" * 60)
    print("CARGA LIBRO IVA DIGITAL - BOT LOGGER")
    print(f"Excel: {EXCEL_PATH} | Workers: {MAX_WORKERS} | Timeout: {TIMEOUT}s | Retries: {MAX_RETRIES}")
    print("=" * 60)

    asyncio.run(main_batch(excel_path=EXCEL_PATH, max_workers=MAX_WORKERS, timeout=TIMEOUT, output_informe=OUTPUT))
    print("\nListo.")
