import asyncio
import json
import os
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Optional

import openpyxl
from dotenv import load_dotenv
from playwright.async_api import (
    Browser,
    BrowserContext,
    Locator,
    Page,
    Playwright,
    TimeoutError as PlaywrightTimeoutError,
    async_playwright,
)

from shared.arca_login import launch_arca_browser_context, login_arca, open_arca_service, wait_page_ready

BOT_NAME = "carga_libro_iva"
TEMP_DIR = Path("./temp")
LOGS_ERROR_DIR = Path("./Logs_Error_Individuales")
INVALID_FILENAME_RE = re.compile(r'[\\/:*?"<>|]+')
PORTAL_IVA_URL_PATTERN = re.compile(r"(siapweb\.cloud\.afip\.gob\.ar/iva|liva\.afip\.gob\.ar)", re.I)
PORTAL_IVA_DIRECT_URL = "https://siapweb.cloud.afip.gob.ar/iva"
PORTAL_IVA_INIT_URL = "https://siapweb.cloud.afip.gob.ar/iva/#/init"
TEMP_DIR.mkdir(exist_ok=True)
LOGS_ERROR_DIR.mkdir(exist_ok=True)

HEADLESS_MODE = os.getenv("CARGA_LIBRO_IVA_HEADLESS", "true").strip().lower() in ("true", "1", "yes")
LIVA_BASE = "https://liva.afip.gob.ar/liva/jsp"

VALORES_PROCESAR = {"si", "sí", "yes", "true", "1"}

# --- Timeouts ---
SHORT = 1_500
MEDIUM = 3_000
LONG = 5_000
CLICK_TO = 10_000
NAV_TO = 15_000
UPLOAD_TO = 60_000
MAX_RETRIES = 3


def _format_cuit(cuit: str) -> str:
    digits = re.sub(r"\D", "", cuit or "")
    if len(digits) == 11:
        return f"{digits[:2]}-{digits[2:10]}-{digits[10:]}"
    return cuit


def _sanitize(value: str) -> str:
    return INVALID_FILENAME_RE.sub("_", (value or "").strip())


def _normalize_period(periodo: str) -> str:
    clean = (periodo or "").strip()
    match = re.fullmatch(r"(\d{2})/(\d{4})", clean)
    if match:
        return f"{match.group(2)}{match.group(1)}"
    match = re.fullmatch(r"(\d{4})(\d{2})", clean)
    if match:
        return clean
    match = re.fullmatch(r"(\d{2})-(\d{4})", clean)
    if match:
        return f"{match.group(2)}{match.group(1)}"
    return re.sub(r"\D", "", clean)


def _looks_like_blank_html(content: str) -> bool:
    compact = re.sub(r"\s+", "", content or "").lower()
    if not compact:
        return True
    if compact in {"<html><head></head><body></body></html>", "<html><head></head><body></body></html>"}:
        return True
    return len(compact) < 120


async def _is_blank_page(page: Page) -> bool:
    try:
        content = await page.content()
    except Exception:
        return True
    return _looks_like_blank_html(content)


def _portal_iva_init_url_with_key() -> str:
    token = format(int(datetime.now().timestamp() * 1000), "x")[-6:]
    return f"{PORTAL_IVA_INIT_URL}?_k={token}"


@dataclass
class ContribuyenteRow:
    nro: str
    cliente: str
    cuit_afip: str
    cuit_en_pagina: str
    clave: str
    cf: str
    tipo: str
    operaciones_ng_e: str
    periodo_afip: str
    ubicacion_txt: str
    txt_liv_c: str
    txt_liv_a: str
    txt_lic_c: str
    txt_lic_a: str
    directorio_iva_simple: str
    iva_simple: str

    @property
    def periodo_normalizado(self) -> str:
        return _normalize_period(self.periodo_afip)

    @property
    def cuit_formateado(self) -> str:
        return _format_cuit(self.cuit_en_pagina)

    @property
    def es_ng_e(self) -> bool:
        val = str(self.operaciones_ng_e or "").strip().upper()
        return val in ("NG/E", "NGE", "SI", "TRUE", "1")

    @property
    def prorrateo_tipo(self) -> Optional[str]:
        val = str(self.cf or "").strip().lower()
        if "global" in val:
            return "global"
        if "directa" in val or "comprobante" in val:
            return "directa"
        if "ambos" in val:
            return "ambos"
        return None


@dataclass
class ResultadoCarga:
    nro: str = ""
    cliente: str = ""
    cuit: str = ""
    periodo: str = ""
    venta_cbte: str = ""
    venta_alicuota: str = ""
    compra_cbte: str = ""
    compra_alicuota: str = ""
    cf_status: str = ""
    cf_restitucion: str = ""
    df_status: str = ""
    df_restitucion: str = ""
    estado_general: str = "PENDIENTE"
    error_msg: str = ""


# --- Column mapping ---
_EXCEL_COLUMNS = [
    ("Procesar", "Procesar"),
    ("Nro", "Nro"),
    ("Cliente", "Cliente"),
    ("CUIT AFIP", "CUIT AFIP"),
    ("CUIT en pagina", "CUIT en pagina"),
    ("CLAVE", "CLAVE"),
    ("CF", "CF"),
    ("Tipo", "Tipo"),
    ("Operaciones NG o E", "Operaciones NG o E"),
    ("Periodo AFIP", "Periodo AFIP"),
    ("Ubicacion TXT", "Ubicacion TXT"),
    ("TXT LIV C", "TXT LIV C"),
    ("TXT LIV A", "TXT LIV A"),
    ("TXT LIC C", "TXT LIC C"),
    ("TXT LIC A", "TXT LIC A"),
    ("Directorio IVA Simple", "Directorio IVA Simple"),
    ("IVA Simple", "IVA Simple"),
]


def _build_col_idx(headers: list[str]) -> dict[str, int]:
    norm_headers = [re.sub(r"[áéíóú]", lambda m: {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u"}[m.group(0)], h.lower().strip()) for h in headers]
    col_idx = {}
    for key, expected in _EXCEL_COLUMNS:
        norm_key = re.sub(r"[áéíóú]", lambda m: {"á": "a", "é": "e", "í": "i", "ó": "o", "ú": "u"}[m.group(0)], key.lower().strip())
        try:
            idx = norm_headers.index(norm_key)
        except ValueError:
            for i, nh in enumerate(norm_headers):
                if norm_key in nh or nh in norm_key:
                    idx = i
                    break
            else:
                continue
        col_idx[key] = idx
    return col_idx


def _str_val(row, col_idx: dict[str, int], key: str) -> str:
    idx = col_idx.get(key)
    if idx is None:
        return ""
    return str(row[idx] or "").strip()


def leer_contribuyentes(excel_path: str) -> list[ContribuyenteRow]:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Listado"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [str(cell.value or "").strip() for cell in ws[1]]
    col_idx = _build_col_idx(headers)

    contribuyentes: list[ContribuyenteRow] = []
    for row in rows:
        procesar = _str_val(row, col_idx, "Procesar").lower()
        tipo = _str_val(row, col_idx, "Tipo").upper()
        if procesar not in VALORES_PROCESAR or tipo != "RI":
            continue

        c = ContribuyenteRow(
            nro=_str_val(row, col_idx, "Nro"),
            cliente=_str_val(row, col_idx, "Cliente"),
            cuit_afip=_str_val(row, col_idx, "CUIT AFIP"),
            cuit_en_pagina=_str_val(row, col_idx, "CUIT en pagina"),
            clave=_str_val(row, col_idx, "CLAVE"),
            cf=_str_val(row, col_idx, "CF"),
            tipo=_str_val(row, col_idx, "Tipo"),
            operaciones_ng_e=_str_val(row, col_idx, "Operaciones NG o E"),
            periodo_afip=_str_val(row, col_idx, "Periodo AFIP"),
            ubicacion_txt=_str_val(row, col_idx, "Ubicacion TXT"),
            txt_liv_c=_str_val(row, col_idx, "TXT LIV C"),
            txt_liv_a=_str_val(row, col_idx, "TXT LIV A"),
            txt_lic_c=_str_val(row, col_idx, "TXT LIC C"),
            txt_lic_a=_str_val(row, col_idx, "TXT LIC A"),
            directorio_iva_simple=_str_val(row, col_idx, "Directorio IVA Simple"),
            iva_simple=_str_val(row, col_idx, "IVA Simple"),
        )
        contribuyentes.append(c)

    wb.close()
    return contribuyentes


# --- Helper functions ---
async def try_click(locator: Locator, timeout: int = CLICK_TO) -> bool:
    try:
        await locator.first.click(timeout=timeout)
        return True
    except Exception:
        return False


async def click_first(locators: list[Locator], timeout: int = CLICK_TO) -> bool:
    for locator in locators:
        try:
            if await locator.count() == 0:
                continue
            if await try_click(locator, timeout=timeout):
                return True
        except Exception:
            continue
    return False


async def is_visible(locator: Locator, timeout: int = SHORT) -> bool:
    try:
        if await locator.count() == 0:
            return False
        await locator.first.wait_for(state="visible", timeout=timeout)
        return True
    except Exception:
        return False


async def esperar_procesando(page: Page) -> None:
    try:
        procesando = page.get_by_text("Procesando", exact=False)
        if await is_visible(procesando, timeout=2_000):
            try:
                await procesando.first.wait_for(state="hidden", timeout=UPLOAD_TO)
            except Exception:
                pass
            try:
                await page.wait_for_timeout(SHORT)
            except Exception:
                pass
    except Exception:
        pass


async def ensure_session_active(page: Page) -> None:
    checks = [
        page.get_by_text("usuario no está logueado", exact=False),
        page.get_by_text("sesión ha expirado", exact=False),
        page.get_by_role("button", name=re.compile(r"IR A LA P[ÁA]GINA DE AFIP", re.I)),
    ]
    for locator in checks:
        if await is_visible(locator, timeout=500):
            raise RuntimeError("La sesión de Portal IVA expiró")


async def return_to_home(page: Page) -> None:
    await ensure_session_active(page)
    await esperar_procesando(page)
    cur_url = getattr(page, "url", "")

    if "menuPresentacion" in cur_url:
        return

    if "liva.afip.gob.ar" in cur_url:
        try:
            await page.goto(f"{LIVA_BASE}/menuPresentacion.do", wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
            await esperar_procesando(page)
            if "menuPresentacion" in getattr(page, "url", ""):
                return
        except Exception:
            pass

    for locators in (
        [page.locator("a[href*='menuPresentacion.do']"), page.locator("#lnkInicio")],
        [page.get_by_role("link", name=re.compile(r"Inicio", re.I)), page.get_by_text("Inicio", exact=False)],
        [page.get_by_role("link", name=re.compile(r"Libro IVA - Borrador", re.I)), page.get_by_text("Libro IVA - Borrador", exact=False)],
        [page.get_by_role("button", name=re.compile(r"Volver", re.I)), page.get_by_text(re.compile(r"Volver", re.I)), page.locator("#btnVolver")],
    ):
        if await click_first(locators, timeout=LONG):
            await page.wait_for_timeout(MEDIUM)
            await esperar_procesando(page)
            return

    try:
        await page.goto(f"{LIVA_BASE}/menuPresentacion.do", wait_until="domcontentloaded", timeout=NAV_TO)
        await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page)
    except Exception:
        pass


# --- Login & Portal ---
async def login_y_abrir_portal(
    playwright: Playwright,
    contrib: ContribuyenteRow,
) -> tuple[Browser, BrowserContext, Page]:
    browser, context, page = await launch_arca_browser_context(
        playwright, headless=HEADLESS_MODE,
        context_kwargs={"accept_downloads": True},
    )
    page.set_default_timeout(30_000)

    login_errors: list[str] = []
    result = await login_arca(page, contrib.cuit_afip, contrib.clave, error_logs=login_errors, fail_on_captcha=True)
    for line in login_errors:
        pass
    if result.captcha_detected:
        raise RuntimeError("Captcha detectado")
    if result.password_change_required:
        raise RuntimeError("Se requiere cambiar la clave fiscal")
    if not result.success:
        raise RuntimeError(login_errors[-1] if login_errors else "No se pudo iniciar sesión en ARCA")

    open_errors: list[str] = []
    portal_page = await open_arca_service(
        page, service_name="Portal IVA",
        service_matchers=(re.compile(r"PORTAL\s+IVA", re.I), "PORTAL IVA Presentación del"),
        error_logs=open_errors, click_ver_todos=True,
        service_unavailable_message="No se encuentra habilitado el servicio de Portal IVA",
    )
    for line in open_errors:
        pass
    if portal_page is None:
        raise RuntimeError(open_errors[-1] if open_errors else "No se encuentra habilitado el servicio de Portal IVA")
    portal_page.set_default_timeout(30_000)

    try:
        await wait_page_ready(portal_page, timeout_ms=20_000)
        await portal_page.wait_for_url(PORTAL_IVA_URL_PATTERN, timeout=20_000)
    except PlaywrightTimeoutError:
        await portal_page.goto(PORTAL_IVA_DIRECT_URL, wait_until="domcontentloaded", timeout=20_000)
        await wait_page_ready(portal_page, timeout_ms=20_000)
        await portal_page.wait_for_url(PORTAL_IVA_URL_PATTERN, timeout=20_000)

    if await _is_blank_page(portal_page):
        await portal_page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
        await wait_page_ready(portal_page, timeout_ms=20_000)

    await portal_page.wait_for_timeout(LONG)
    return browser, context, portal_page


# --- Representado ---
async def seleccionar_representado(page: Page, contrib: ContribuyenteRow) -> None:
    formato = contrib.cuit_formateado
    target_digits = re.sub(r"\D", "", contrib.cuit_en_pagina or "")
    login_digits = re.sub(r"\D", "", contrib.cuit_afip or "")

    async def _is_target() -> bool:
        if await is_visible(page.get_by_text(re.compile(rf"Representando a:.*{re.escape(formato)}", re.I)), timeout=1_200):
            return True
        if target_digits == login_digits and await is_visible(page.get_by_text(formato, exact=False), timeout=1_200):
            return True
        return False

    await ensure_session_active(page)
    try:
        await wait_page_ready(page, timeout_ms=NAV_TO)
    except Exception:
        pass
    await page.wait_for_timeout(SHORT)

    for _ in range(2):
        if not await _is_blank_page(page):
            break
        await page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
        await wait_page_ready(page, timeout_ms=20_000)
        await page.wait_for_timeout(MEDIUM)

    if await _is_blank_page(page):
        raise RuntimeError("No se pudo cargar Portal IVA")
    if await _is_target():
        return

    await click_first([
        page.get_by_title("cambio relación"),
        page.locator("[title*='cambio relación' i]"),
        page.get_by_role("link", name=""),
    ], timeout=8_000)
    await page.wait_for_timeout(3_000)

    representado_link = page.locator('a[title="Representar a..."]').filter(
        has_text=re.compile(rf"{re.escape(contrib.cliente)}|{re.escape(formato)}", re.I)
    )
    selected = False
    if await representado_link.count() > 0:
        await representado_link.first.click(timeout=8_000)
        selected = True
        await page.wait_for_timeout(3_000)

    current_html = await page.content()
    for _ in range(2):
        if not _looks_like_blank_html(current_html):
            break
        await page.goto(_portal_iva_init_url_with_key(), wait_until="domcontentloaded", timeout=20_000)
        await wait_page_ready(page, timeout_ms=20_000)
        await page.wait_for_timeout(MEDIUM)
        current_html = await page.content()

    if await _is_target():
        return

    available_cuit_digits = set()
    for raw_cuit in re.findall(r"\b\d{2}-\d{8}-\d\b", current_html):
        available_cuit_digits.add(re.sub(r"\D", "", raw_cuit))
    target_present = target_digits in available_cuit_digits

    if not target_present:
        raise RuntimeError("CUIT sin representados válidos con acceso a liquidaciones de IVA")

    raise RuntimeError(f"No se pudo seleccionar el representado {formato}")


# --- LIVA Navigation + Configuration ---
async def navegar_a_liva_y_configurar(page: Page, contrib: ContribuyenteRow) -> None:
    await ensure_session_active(page)
    await esperar_procesando(page)

    if not await is_visible(page.locator("#periodo"), timeout=2_000):
        await click_first([
            page.get_by_role("button", name=re.compile(r"iva\.home\.btn\.nueva", re.I)),
            page.get_by_label("Sin texto (iva.home.btn.nueva"),
            page.get_by_role("button", name=re.compile(r"Ingresar", re.I)),
            page.get_by_text(re.compile(r"Nueva|Ingresar", re.I)),
        ], timeout=NAV_TO)
        await page.wait_for_timeout(MEDIUM)

    if await page.locator("#periodo").count() > 0:
        await page.locator("#periodo").select_option(contrib.periodo_normalizado)
        await page.wait_for_timeout(SHORT)
        await click_first([
            page.get_by_role("button", name=re.compile(r"Validar", re.I)),
            page.get_by_label("Sin texto (iva.btn.home.validar.periodo.alt)"),
        ], timeout=12_000)
        await page.wait_for_timeout(MEDIUM)
        await esperar_procesando(page)

    if "liva" not in getattr(page, "url", ""):
        await click_first([
            page.get_by_role("button", name=re.compile(r"Registraci[oó]n y declaraci[oó]n", re.I)),
            page.get_by_role("button", name=re.compile(r"Registraci[oó]n", re.I)),
            page.get_by_role("button", name=re.compile(r"iva\.btn\.home\.liva", re.I)),
            page.get_by_label("Sin texto (iva.btn.home.liva."),
            page.get_by_role("button", name=re.compile(r"Libro\s+IVA\s+Digital", re.I)),
            page.locator("button").filter(has_text=re.compile(r"Registraci[oó]n|Libro IVA|Declaraci[oó]n", re.I)),
        ], timeout=NAV_TO)
        await page.wait_for_timeout(LONG)
        await wait_page_ready(page, timeout_ms=20_000)
        await esperar_procesando(page)

    if "menuPresentacion" in getattr(page, "url", ""):
        return

    if "vistaPreviaSimplificada" in getattr(page, "url", ""):
        await click_first([
            page.locator("#btnModificar"),
            page.get_by_role("link", name=re.compile(r"Modificar", re.I)),
            page.get_by_role("button", name=re.compile(r"Modificar", re.I)),
        ], timeout=12_000)
        await page.wait_for_timeout(MEDIUM)
        await esperar_procesando(page)
        if "menuPresentacion" in getattr(page, "url", ""):
            return

    if await is_visible(page.locator("#btnVistaPrevia"), timeout=3_000):
        return

    if contrib.es_ng_e:
        await click_first([
            page.locator("label").filter(has_text=re.compile(r"No Gravado", re.I)).locator("span"),
            page.locator("#operacionesNoGravadasExentas"),
        ], timeout=LONG)
        await page.wait_for_timeout(SHORT)
        prorrateo = contrib.prorrateo_tipo
        if prorrateo:
            if await page.locator("#tipoProrrateo").count() > 0:
                try:
                    await page.locator("#tipoProrrateo").select_option(
                        "3" if prorrateo == "ambos" else ("2" if prorrateo == "global" else "1"))
                except Exception:
                    pass
            else:
                await click_first([page.get_by_role("button", name=re.compile(r"prorrateo", re.I))], timeout=LONG)
                await page.wait_for_timeout(SHORT)
                opts = {
                    "global": ["Con prorrateo global"],
                    "directa": ["Con prorrateo por asignación directa"],
                    "ambos": ["Con prorrateo por asignación directa y global"],
                }
                await click_first([
                    page.get_by_role("option", name=re.compile(rf"{opts[prorrateo][0]}", re.I)),
                ], timeout=LONG)
    else:
        await click_first([
            page.locator("label").filter(has_text=re.compile(r"No realizo ninguna", re.I)).locator("span"),
            page.locator("#ningunaOperacion"),
        ], timeout=LONG)

    await page.wait_for_timeout(SHORT)
    guardar = page.locator("#btnGuardar").first
    if await guardar.count() > 0:
        await guardar.click(timeout=CLICK_TO)
        await page.wait_for_timeout(MEDIUM)


# --- Libros: abrir, eliminar, importar ---
async def abrir_libro(page: Page, book: str) -> None:
    href = "verVentas.do" if book == "ventas" else "verCompras.do"
    label = "Libro Ventas" if book == "ventas" else "Libro Compras"

    await ensure_session_active(page)
    await esperar_procesando(page)
    clicked = await click_first([
        page.locator(f"a[href*='{href}']"),
        page.get_by_role("link", name=re.compile(rf"{label}\s+Registr", re.I)),
        page.get_by_role("link", name=re.compile(label, re.I)),
        page.locator(f"#{'btnLibroVentas' if book == 'ventas' else 'btnLibroCompras'}"),
    ], timeout=NAV_TO)
    if not clicked:
        raise RuntimeError(f"No se pudo abrir {label}")
    await page.wait_for_timeout(MEDIUM)


async def eliminar_todo(page: Page) -> bool:
    await ensure_session_active(page)
    await page.wait_for_timeout(MEDIUM)
    deleted = await click_first([
        page.get_by_role("button", name=re.compile(r"Eliminar\s+Todos?\.\.\.?", re.I)),
        page.locator("#btnEliminarTodo"),
    ], timeout=8_000)
    if not deleted:
        return False
    await page.wait_for_timeout(SHORT)
    confirmed = await click_first([
        page.locator(".bootstrap-dialog .btn-danger").filter(has_text=re.compile(r"Eliminar", re.I)),
        page.locator(".modal.in .btn-danger"),
        page.get_by_role("button", name=re.compile(r"Eliminar", re.I)),
    ], timeout=LONG)
    if not confirmed:
        await click_first([page.locator(".bootstrap-dialog .close")], timeout=MEDIUM)
        return False
    await page.wait_for_timeout(MEDIUM)
    return True


async def importar_txts(page: Page, book: str, cbte_path: str, alicuota_path: str) -> None:
    comp_label = "Archivo Ventas *" if book == "ventas" else "Archivo Compras *"
    alic_label = "Archivo IVA Ventas *" if book == "ventas" else "Archivo IVA Compras *"

    await ensure_session_active(page)
    opened = await click_first([
        page.locator("#btnDropdownImportar"),
        page.get_by_role("button", name=re.compile(r"Importar", re.I)),
    ], timeout=CLICK_TO)
    if not opened:
        raise RuntimeError(f"No se encontró Importar para {book}")
    await page.wait_for_timeout(SHORT)

    clicked = await click_first([
        page.get_by_role("link", name=re.compile(r"Importar\s+(?:desde\s+)?Archivos?", re.I)),
        page.locator("#lnkImportarArchivo"),
    ], timeout=12_000)
    if not clicked:
        raise RuntimeError(f"No se encontró 'Importar Archivos' para {book}")
    await page.wait_for_timeout(MEDIUM)

    await page.get_by_label(comp_label).set_input_files(cbte_path)
    await page.get_by_label(alic_label).set_input_files(alicuota_path)

    await page.wait_for_timeout(SHORT)
    await page.locator("#btnImportarArchivosImportar").evaluate("el => el.click()")
    await page.wait_for_timeout(MEDIUM)

    await esperar_procesando(page)
    try:
        await page.get_by_text("Procesada", exact=False).first.wait_for(state="visible", timeout=UPLOAD_TO)
        await click_first([
            page.get_by_text("Procesada", exact=False),
            page.locator(".bootstrap-dialog .close"),
        ], timeout=LONG)
    except Exception:
        pass
    await page.wait_for_timeout(SHORT)
    await click_first([page.locator("#btnTareasCerrar")], timeout=LONG)


# --- Helper: Crédito Fiscal Prorrateo ---
def _calcular_total_cf_csv(csv_path: str) -> str | None:
    if not os.path.isfile(csv_path):
        return None
    try:
        total = 0.0
        with open(csv_path, encoding="utf-8-sig") as f:
            lines = f.readlines()
        if not lines:
            return None
        header = lines[0].strip().split(";")
        col_idx = next((i for i, h in enumerate(header) if "Credito Fiscal Computable" in h), None)
        if col_idx is None:
            return None
        for line in lines[1:]:
            line = line.strip()
            if not line:
                continue
            parts = line.split(";")
            if len(parts) > col_idx:
                val_str = parts[col_idx].strip()
                if val_str:
                    val_str = val_str.replace(".", "").replace(",", ".")
                    total += float(val_str)
        return f"{total:.2f}"
    except Exception:
        return None


async def _handle_credito_fiscal_modal(page: Page, valor: str | None = None) -> bool:
    input_field = page.locator("#credFiscProrrateo")
    if await input_field.count() == 0:
        return False

    if valor is not None:
        await input_field.first.fill(valor)
        await page.wait_for_timeout(SHORT)
        return True

    try:
        body_text = await page.locator("body").inner_text()
    except Exception:
        body_text = ""
    numeros = re.findall(r'(?:\$?\s*)([\d]{1,3}(?:\.\d{3})*(?:,\d{2})?)', body_text)
    if not numeros:
        numeros = re.findall(r'(\d{1,3}(?:\.\d{3})*(?:,\d{2})?)', body_text)
    if numeros:
        raw = numeros[-1]
        limpio = raw.replace(".", "").replace(",", ".")
        await input_field.first.fill(limpio)
        await page.wait_for_timeout(SHORT)
        return True
    return False


async def _ir_a_editar_credito_fiscal_en_compras(page: Page, csv_path: str = "") -> bool:
    valor = _calcular_total_cf_csv(csv_path) if csv_path else None

    await return_to_home(page)
    await esperar_procesando(page)

    cf_url = f"{LIVA_BASE}/verComprasCreditoFiscalComputableGlobal.do"
    try:
        await page.goto(cf_url, wait_until="domcontentloaded", timeout=NAV_TO)
        await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page)
        if await page.locator("#credFiscProrrateo").count() > 0:
            manejado = await _handle_credito_fiscal_modal(page, valor=valor)
            if manejado:
                await click_first([
                    page.locator("button[type='submit']"),
                    page.get_by_role("button", name=re.compile(r"Guardar", re.I)),
                ], timeout=8_000)
                await page.wait_for_timeout(MEDIUM)
                await esperar_procesando(page)
                return True
    except Exception:
        pass

    await abrir_libro(page, "compras")
    await esperar_procesando(page)

    clicked = await click_first([
        page.locator("a, button, span, div").filter(has_text=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
        page.get_by_role("link", name=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
        page.locator("a[href*='verComprasCreditoFiscalComputableGlobal']"),
        page.get_by_text(re.compile(r"CF\s+Computable\s+Global|Crédito\s+Fiscal\s+Computable", re.I), exact=False),
    ], timeout=8_000)
    if not clicked:
        await return_to_home(page)
        await esperar_procesando(page)
        try:
            if await page.locator("#btnAjustes").count() > 0:
                await page.locator("#btnAjustes").first.click(timeout=CLICK_TO)
                await page.wait_for_timeout(MEDIUM)
            if await page.locator("a[href*='menuInfoAdicional']").count() > 0:
                await page.locator("a[href*='menuInfoAdicional']").first.click(timeout=CLICK_TO)
                await page.wait_for_timeout(MEDIUM)
            else:
                await page.goto(f"{LIVA_BASE}/menuInfoAdicional.jsp", wait_until="domcontentloaded", timeout=NAV_TO)
                await wait_page_ready(page, timeout_ms=10_000)
        except Exception:
            await page.goto(f"{LIVA_BASE}/menuInfoAdicional.jsp", wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page)
        clicked = await click_first([
            page.locator("a, button, span, div").filter(has_text=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
            page.get_by_role("link", name=re.compile(r"Crédito\s+Fiscal\s+Computable\s+Global", re.I)),
            page.locator("a[href*='verComprasCreditoFiscalComputableGlobal']"),
            page.get_by_text(re.compile(r"Crédito\s+Fiscal", re.I), exact=False),
            page.locator("a.btn-primary").filter(has_text=re.compile(r"Crédito\s+Fiscal", re.I)),
            page.locator("a[href*='creditoFiscal']"),
        ], timeout=8_000)
    if not clicked:
        return False
    await page.wait_for_timeout(MEDIUM)
    await esperar_procesando(page)
    manejado = await _handle_credito_fiscal_modal(page, valor=valor)

    await click_first([
        page.locator("button[type='submit']"),
        page.get_by_role("button", name=re.compile(r"Guardar", re.I)),
    ], timeout=8_000)
    await page.wait_for_timeout(MEDIUM)
    await esperar_procesando(page)

    try:
        await page.goto(f"{LIVA_BASE}/menuPresentacion.do", wait_until="domcontentloaded", timeout=NAV_TO)
        await wait_page_ready(page, timeout_ms=10_000)
        await esperar_procesando(page)
    except Exception:
        pass

    return manejado


# --- Apertura CSV ---
async def cargar_apertura_seccion(page: Page, seccion: str, csv_path: str) -> str:
    if not os.path.isfile(csv_path):
        return "Sin archivo"
    if os.path.getsize(csv_path) == 0:
        return "Sin movimientos"
    try:
        with open(csv_path, encoding="utf-8-sig") as f:
            non_empty_lines = [line for line in f if line.strip()]
        if len(non_empty_lines) <= 1:
            return "Sin movimientos"
    except Exception:
        pass

    links_map = {
        "CF": re.compile(r"Crédito\s+Fiscal(?!\s+Comp)", re.I),
        "CF Restitucion": re.compile(r"Restituci[oó]n.*Crédito\s+Fiscal", re.I),
        "DF": re.compile(r"Débito\s+Fiscal(?!\s+Comp)", re.I),
        "DF Restitucion": re.compile(r"Restituci[oó]n.*Débito\s+Fiscal", re.I),
    }
    link_pattern = links_map.get(seccion)

    async def _safe_goto(page: Page, url: str) -> bool:
        try:
            await page.goto(url, wait_until="domcontentloaded", timeout=NAV_TO)
            await wait_page_ready(page, timeout_ms=10_000)
            return True
        except Exception:
            return False

    cur = getattr(page, "url", "")
    if "menuInfoAdicional" not in cur:
        await esperar_procesando(page)
        try:
            if await page.locator("#btnAjustes").count() > 0:
                await page.locator("#btnAjustes").first.click(timeout=CLICK_TO)
                await page.wait_for_timeout(MEDIUM)
            if await page.locator("a[href*='menuInfoAdicional']").count() > 0:
                await page.locator("a[href*='menuInfoAdicional']").first.click(timeout=CLICK_TO)
                await page.wait_for_timeout(MEDIUM)
            else:
                await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp")
        except Exception:
            await _safe_goto(page, f"{LIVA_BASE}/menuInfoAdicional.jsp")
        await esperar_procesando(page)

    href_map = {
        "CF": "d=c", "CF Restitucion": "d=rc",
        "DF": "d=d", "DF Restitucion": "d=rd",
    }
    clicked = await click_first([
        page.locator(f".form-group a.btn-primary[href*='{href_map[seccion]}']"),
        page.locator("a.btn-primary").filter(has_text=link_pattern),
        page.get_by_role("link", name=link_pattern),
    ], timeout=NAV_TO)
    if not clicked:
        raise RuntimeError(f"No se encontró link para {seccion}")
    await page.wait_for_timeout(MEDIUM)
    await esperar_procesando(page)

    btn_importar = page.locator("#btnImportarArchivo")
    if await btn_importar.count() > 0:
        try:
            await btn_importar.first.click(timeout=CLICK_TO)
        except Exception:
            await btn_importar.first.evaluate("el => el.click()")
    else:
        await click_first([
            page.locator("button.btn-link").filter(has_text=re.compile(r"Importar", re.I)),
            page.get_by_role("button", name=re.compile(r"Importar.*CSV|Importar.*Archivo", re.I)),
        ], timeout=CLICK_TO)
    await page.wait_for_timeout(MEDIUM)

    file_input = page.locator("input[type='file']")
    if await file_input.count() > 0:
        await file_input.first.set_input_files(csv_path)
    await page.wait_for_timeout(SHORT)

    importar_btn = page.locator("#btnImportarArchivoImportar")
    if await importar_btn.count() > 0:
        try:
            await importar_btn.first.click(timeout=CLICK_TO)
        except Exception:
            try:
                await importar_btn.evaluate("el => el.click()")
            except Exception:
                pass
    await page.wait_for_timeout(MEDIUM)

    try:
        await esperar_procesando(page)
    except Exception:
        pass
    try:
        await page.wait_for_timeout(2_000)
    except Exception:
        pass

    if seccion == "CF":
        try:
            await _handle_credito_fiscal_modal(page)
        except Exception:
            pass

    try:
        await click_first([
            page.locator("button[type='submit']"),
            page.get_by_role("button", name=re.compile(r"Guardar", re.I)),
        ], timeout=8_000)
        await page.wait_for_timeout(MEDIUM)
    except Exception:
        pass
    try:
        await esperar_procesando(page)
    except Exception:
        pass

    try:
        await click_first([
            page.locator("a[href*='menuInfoAdicional']"),
            page.locator("#btnVolver"),
            page.get_by_text(re.compile(r"Volver", re.I)),
            page.locator("a[href*='menuPresentacion']"),
        ], timeout=8_000)
        await page.wait_for_timeout(MEDIUM)
    except Exception:
        pass

    return "OK"


# --- Procesamiento por contribuyente ---
async def procesar_contribuyente(playwright: Playwright, contrib: ContribuyenteRow) -> ResultadoCarga:
    res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado, periodo=contrib.periodo_afip)
    cuit_sanitized = _sanitize(contrib.cuit_afip) or _sanitize(contrib.cuit_en_pagina)
    browser: Optional[Browser] = None
    context: Optional[BrowserContext] = None
    portal_page: Optional[Page] = None

    try:
        browser, context, portal_page = await login_y_abrir_portal(playwright, contrib)
        await seleccionar_representado(portal_page, contrib)
        await navegar_a_liva_y_configurar(portal_page, contrib)

        if "liva" not in getattr(portal_page, "url", ""):
            raise RuntimeError(f"No se pudo acceder al sistema Libro IVA. URL: {getattr(portal_page, 'url', '')}")

        # --- TXTs Ventas ---
        txt_liv_c = os.path.join(contrib.ubicacion_txt, contrib.txt_liv_c) if contrib.ubicacion_txt and contrib.txt_liv_c else ""
        txt_liv_a = os.path.join(contrib.ubicacion_txt, contrib.txt_liv_a) if contrib.ubicacion_txt and contrib.txt_liv_a else ""

        if contrib.txt_liv_c and contrib.txt_liv_a and os.path.isfile(txt_liv_c) and os.path.isfile(txt_liv_a):
            if os.path.getsize(txt_liv_c) == 0 and os.path.getsize(txt_liv_a) == 0:
                res.venta_cbte = "Sin movimientos"
                res.venta_alicuota = "Sin movimientos"
            else:
                try:
                    await abrir_libro(portal_page, "ventas")
                    await eliminar_todo(portal_page)
                    await importar_txts(portal_page, "ventas", txt_liv_c, txt_liv_a)
                    res.venta_cbte = "OK"
                    res.venta_alicuota = "OK"
                except Exception as e:
                    res.venta_cbte = "ERROR"
                    res.venta_alicuota = "ERROR"
        else:
            res.venta_cbte = "Sin archivo"
            res.venta_alicuota = "Sin archivo"

        await return_to_home(portal_page)

        # --- TXTs Compras ---
        txt_lic_c = os.path.join(contrib.ubicacion_txt, contrib.txt_lic_c) if contrib.ubicacion_txt and contrib.txt_lic_c else ""
        txt_lic_a = os.path.join(contrib.ubicacion_txt, contrib.txt_lic_a) if contrib.ubicacion_txt and contrib.txt_lic_a else ""

        if contrib.txt_lic_c and contrib.txt_lic_a and os.path.isfile(txt_lic_c) and os.path.isfile(txt_lic_a):
            if os.path.getsize(txt_lic_c) == 0 and os.path.getsize(txt_lic_a) == 0:
                res.compra_cbte = "Sin movimientos"
                res.compra_alicuota = "Sin movimientos"
            else:
                try:
                    await abrir_libro(portal_page, "compras")
                    await eliminar_todo(portal_page)
                    await importar_txts(portal_page, "compras", txt_lic_c, txt_lic_a)
                    res.compra_cbte = "OK"
                    res.compra_alicuota = "OK"
                except Exception as e:
                    res.compra_cbte = "ERROR"
                    res.compra_alicuota = "ERROR"
        else:
            res.compra_cbte = "Sin archivo"
            res.compra_alicuota = "Sin archivo"

        await return_to_home(portal_page)

        # --- Aperturas CSV ---
        await esperar_procesando(portal_page)
        for seccion_key, res_attr in [
            ("CF", "cf_status"), ("CF Restitucion", "cf_restitucion"),
            ("DF", "df_status"), ("DF Restitucion", "df_restitucion"),
        ]:
            csv_filename = contrib.iva_simple.replace("(XXXX)", seccion_key)
            csv_path = os.path.join(contrib.directorio_iva_simple, csv_filename) if contrib.directorio_iva_simple else ""
            try:
                status = await cargar_apertura_seccion(portal_page, seccion_key, csv_path)
                setattr(res, res_attr, status)
                if seccion_key == "CF" and status == "OK" and contrib.prorrateo_tipo == "global":
                    try:
                        await _ir_a_editar_credito_fiscal_en_compras(portal_page, csv_path=csv_path)
                    except Exception:
                        pass
                if status in ("OK", "ERROR"):
                    try:
                        await return_to_home(portal_page)
                    except Exception:
                        pass
            except Exception as e:
                setattr(res, res_attr, "ERROR")
                res.error_msg = str(e)
                try:
                    await return_to_home(portal_page)
                except Exception:
                    pass

        # --- Estado general ---
        estados = [res.venta_cbte, res.venta_alicuota, res.compra_cbte, res.compra_alicuota,
                   res.cf_status, res.cf_restitucion, res.df_status, res.df_restitucion]
        if any(e == "ERROR" for e in estados if e):
            res.estado_general = "ERROR"
        else:
            res.estado_general = "COMPLETO"

    except Exception as e:
        res.estado_general = "ERROR"
        res.error_msg = str(e)
        error_log = {"nro": contrib.nro, "cliente": contrib.cliente, "cuit_afip": contrib.cuit_afip,
                     "periodo": contrib.periodo_afip, "error": str(e),
                     "fecha_ejecucion": datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
        with open(LOGS_ERROR_DIR / f"log_{cuit_sanitized}_{contrib.nro}_{contrib.periodo_normalizado}.json", "w") as f:
            json.dump(error_log, f, ensure_ascii=False, indent=4)

    finally:
        if context:
            try:
                await context.close()
            except Exception:
                pass
        if browser:
            try:
                await browser.close()
            except Exception:
                pass
        print(f"[{res.estado_general}] {contrib.cliente} ({res.cuit}) - {contrib.periodo_afip}")
        if res.error_msg:
            print(f"   Error: {res.error_msg[:200]}")

    return res


# --- Worker con retry ---
async def worker(semaphore: asyncio.Semaphore, contrib: ContribuyenteRow,
                  resultados: list[ResultadoCarga], timeout: int = 300) -> None:
    async with semaphore:
        for intento in range(1, MAX_RETRIES + 1):
            print(f"[INICIANDO intento {intento}/{MAX_RETRIES}] {contrib.cliente} ({contrib.cuit_formateado}) - {contrib.periodo_afip}")
            try:
                async with async_playwright() as playwright:
                    res = await asyncio.wait_for(procesar_contribuyente(playwright, contrib), timeout=timeout)
                if res.estado_general == "ERROR" and intento < MAX_RETRIES:
                    print(f"[REINTENTO {intento}] {contrib.cliente} - error: {res.error_msg[:100]}")
                    await asyncio.sleep(3)
                    continue
                resultados.append(res)
                return
            except asyncio.TimeoutError:
                if intento < MAX_RETRIES:
                    print(f"[TIMEOUT intento {intento}] {contrib.cliente} - reintentando...")
                    await asyncio.sleep(3)
                    continue
                print(f"[TIMEOUT] {contrib.cliente}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="TIMEOUT", error_msg=f"Timeout {timeout}s tras {MAX_RETRIES} intentos")
                resultados.append(res)
                return
            except Exception as e:
                if intento < MAX_RETRIES:
                    print(f"[REINTENTO {intento}] {contrib.cliente} - {e}")
                    await asyncio.sleep(3)
                    continue
                print(f"[EXCEPCION] {contrib.cliente} - {e}")
                res = ResultadoCarga(nro=contrib.nro, cliente=contrib.cliente, cuit=contrib.cuit_formateado,
                                     periodo=contrib.periodo_afip, estado_general="ERROR", error_msg=str(e))
                resultados.append(res)
                return


# --- Informe ---
def generar_informe(resultados: list[ResultadoCarga], output_path: str) -> None:
    import pandas as pd
    headers = ["#", "Contribuyente", "CUIT", "Período",
               "Venta CBTE", "Venta Alic.", "Compra CBTE", "Compra Alic.",
               "Apertura CF", "Apertura CF Rest.", "Apertura DF", "Apertura DF Rest.", "Estado"]
    rows = []
    errores = []
    for i, r in enumerate(resultados, 1):
        rows.append([i, r.cliente, r.cuit, r.periodo,
                     r.venta_cbte, r.venta_alicuota, r.compra_cbte, r.compra_alicuota,
                     r.cf_status, r.cf_restitucion, r.df_status, r.df_restitucion, r.estado_general])
        if r.estado_general == "ERROR" and r.error_msg:
            errores.append([r.cliente, r.cuit, "", r.error_msg, ""])
    df_resumen = pd.DataFrame(rows, columns=headers)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_resumen.to_excel(writer, sheet_name="Resumen", index=False)
        if errores:
            pd.DataFrame(errores, columns=["Contribuyente", "CUIT", "Paso", "Error", "Fecha"]).to_excel(
                writer, sheet_name="Detalle Errores", index=False)
    print(f"\nInforme guardado en: {output_path}")


# --- Main ---
async def main_batch(excel_path: str, max_workers: int = 4, timeout: int = 300,
                     output_informe: str = "informe_carga_libro_iva.xlsx") -> list[ResultadoCarga]:
    contribuyentes = leer_contribuyentes(excel_path)
    if not contribuyentes:
        print("No se encontraron contribuyentes para procesar.")
        return []

    print(f"Contribuyentes a procesar: {len(contribuyentes)}")
    for c in contribuyentes:
        print(f"  - [{c.nro}] {c.cliente} ({_format_cuit(c.cuit_en_pagina)}) periodo={c.periodo_afip} NG/E={c.es_ng_e} CF={c.prorrateo_tipo}")

    semaphore = asyncio.Semaphore(max_workers)
    resultados: list[ResultadoCarga] = []
    tasks = [asyncio.create_task(worker(semaphore, c, resultados, timeout)) for c in contribuyentes]
    await asyncio.gather(*tasks)
    resultados.sort(key=lambda r: int(r.nro) if r.nro.isdigit() else 999999)
    generar_informe(resultados, output_informe)

    completos = sum(1 for r in resultados if r.estado_general == "COMPLETO")
    parciales = sum(1 for r in resultados if r.estado_general.startswith("PARCIAL"))
    errores = sum(1 for r in resultados if r.estado_general in ("ERROR", "TIMEOUT"))
    print(f"\nResumen: {completos} completos, {parciales} parciales, {errores} errores/timeouts")
    return resultados


if __name__ == "__main__":
    load_dotenv()
    EXCEL_PATH = os.getenv("CARGA_LIBRO_IVA_EXCEL", "Lista libros SOS.xlsx").strip()
    MAX_WORKERS = int(os.getenv("CARGA_LIBRO_IVA_MAX_WORKERS", "1"))
    TIMEOUT = int(os.getenv("CARGA_LIBRO_IVA_TIMEOUT", "600"))
    OUTPUT = os.getenv("CARGA_LIBRO_IVA_OUTPUT", "informe_carga_libro_iva.xlsx").strip()

    print("=" * 60)
    print("CARGA LIBRO IVA DIGITAL")
    print(f"Excel: {EXCEL_PATH} | Workers: {MAX_WORKERS} | Timeout: {TIMEOUT}s")
    print("=" * 60)

    asyncio.run(main_batch(excel_path=EXCEL_PATH, max_workers=MAX_WORKERS, timeout=TIMEOUT, output_informe=OUTPUT))
    print("\nListo.")
